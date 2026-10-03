"""
build_example_results_v0_1.py - worked-example workbook for the ad hoc Table 2 rebuild (v0.3).

Walks one example per level of calculation (parameters -> CPAT inputs -> one product -> CBAM block ->
national Table 2 metric -> bundle variants), all as live formulas, with the v0.3 workbook value as a check and
a column saying where the main version (kernel v0.15) differs. Inputs are copied from
AdHocCalculations_Rebuild_v0.3.xlsx (Inputs) and cpat_outputs_egypt_2022_2041.csv (2030).
Output: ExampleResults_v0.1.xlsx (run recalc via Excel to cache values).
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "ExampleResults_v0.1.xlsx")
CSV_SRC = os.path.join(HERE, "cpat_outputs_egypt_2022_2041.csv")
YR = 2030

GREEN = PatternFill("solid", fgColor="EBF1DE")
TITLE = PatternFill("solid", fgColor="00B050")
BAND = PatternFill("solid", fgColor="92D050")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def title(ws, text, width):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    for c in range(1, width + 1):
        ws.cell(1, c).fill = TITLE


# --------------------------------------------------------------------------------------------------------------
# Inputs
# --------------------------------------------------------------------------------------------------------------
wsI = wb.active
wsI.title = "Inputs"
title(wsI, "Inputs (copied from AdHoc rebuild v0.3 Inputs and CPAT Outputs, 2030)", 12)
params = [
    ("Yr", "Reporting year", YR, ""),
    ("BaseYr", "Base year of product activity Q0", 2024, ""),
    ("Tau", "Carbon price tau (CPAT cptraj.2, EG1 and EG3, 2030)", 20, "$/t"),
    ("SigmaEff", "Fund shadow price sigma", 20, "$/t"),
    ("EpsU", "eps_U", -0.5, ""),
    ("EpsF", "eps_F", -0.5, ""),
    ("EpsQ", "Output price elasticity eps_Q", -0.5, ""),
    ("PEU", "EU ETS price P_EU", 100, "$/t"),
    ("PStar", "IPCC anchor price P* (USD2019 100 deflated to USD2024)", 122, "$/t"),
    ("CBF", "CBAM phase-in factor (NOPHASE convention = 1)", 1, ""),
    ("Ssel", "Deduction factor S (NOPHASE = 1)", 1, ""),
]
wsI["A3"] = "A. Switches and parameters"
wsI["A3"].font = BOLD
r = 4
for code, label, v, unit in params:
    wsI.cell(r, 1, code)
    wsI.cell(r, 2, label)
    c = wsI.cell(r, 3, v)
    c.fill = GREEN
    wsI.cell(r, 4, unit)
    name(code, "Inputs!$C$%d" % r)
    r += 1

# CPAT values
rows = list(csv.reader(open(CSV_SRC, encoding="utf-8")))
ycol = rows[0].index(str(YR))
CP = {(x[0], x[1]): float(x[ycol] or 0) for x in rows[1:] if x and x[0]}
cpat_codes = [
    ("GHG0", "egy.mit.ghg.tot.inc.1", "Total GHG incl. LULUCF, baseline", "Mt"),
    ("GHG1", "egy.mit.ghg.tot.inc.2", "Total GHG, policy", "Mt"),
    ("IPPU0", "egy.mit.ghg.ipr.tot.1", "IPPU GHG, baseline", "Mt"),
    ("IPPU1", "egy.mit.ghg.ipr.tot.2", "IPPU GHG, policy", "Mt"),
    ("ENR0", "egy.mit.co2.enr.tot.1", "Energy CO2, baseline", "Mt"),
    ("ENR1", "egy.mit.co2.enr.tot.2", "Energy CO2, policy", "Mt"),
    ("IND0", "egy.mit.co2.ind.1", "Industry energy CO2, baseline", "Mt"),
    ("IND1", "egy.mit.co2.ind.2", "Industry energy CO2, policy", "Mt"),
    ("CPT", "egy.mit.cptraj.2", "Carbon price trajectory", "$/t"),
    ("EFFCPT", "egy.mit.eff.cptraj.2", "Effective (coverage-weighted) carbon price", "$/t"),
    ("REVCTAX", None, "Carbon-tax receipts, six fuels (coa+die+gso+lpk+nga+oil .usd.2)", "$bn"),
    ("DEATHS", None, "Deaths avoided (ada.2464 + ada.65 + ada.u24)", "#"),
    ("MORT", "egy.air.mort", "Baseline deaths", "#"),
    ("DREV", None, "Delta net new revenue (rev.new.usd.2 - .1)", "$bn"),
]
r += 1
wsI.cell(r, 1, "B. CPAT values, %d" % YR).font = BOLD
r += 1
for j, h in enumerate(["Code", "CPAT code", "Label", "EG1", "EG3", "Unit"]):
    wsI.cell(r, 1 + j, h).font = BOLD
r += 1
for code, cc, label, unit in cpat_codes:
    wsI.cell(r, 1, code)
    wsI.cell(r, 2, cc or "(sum)")
    wsI.cell(r, 3, label)
    for k, sc in enumerate(["EG1", "EG3"]):
        if code == "REVCTAX":
            v = sum(CP[("egy.mit.rev.new.%s.usd.2" % f, sc)] for f in ["coa", "die", "gso", "lpk", "nga", "oil"])
        elif code == "DEATHS":
            v = sum(CP[("egy.air.%s" % f, sc)] for f in ["ada.2464", "ada.65", "ada.u24"])
        elif code == "DREV":
            v = CP[("egy.mit.rev.new.usd.2", sc)] - CP[("egy.mit.rev.new.usd.1", sc)]
        else:
            v = CP[(cc, sc)]
        c = wsI.cell(r, 4 + k, round(v, 6))
        c.fill = GREEN
        name("%s_%s" % (code, sc), "Inputs!$%s$%d" % ("DE"[k], r))
    wsI.cell(r, 6, unit)
    r += 1

# Products
PROD = [
    ("DRI-EAF steel", 5100, 0.044, 580, 750, 0.57099077, 0.03790102, 0.3433),
    ("Scrap-EAF steel", 4800, 0.044, 300, 680, 0.04488, 0.04398115, 0.181),
    ("BF-BOF steel", 0, 0, 0, 620, 1.42498027, 0.05276542, 0.2803),
    ("Clinker", 50000, 0.033, 800, 110, 0.3135685, 0.53702616, 0.3352),
    ("Ammonia", 1785, 0.022, 120, 450, 1.97472, 0, 0.2974),
    ("Urea", 2800, 0.022, 1600, 380, 0.1122, 0, 0.1158),
    ("Ammonium nitrate", 600, 0.022, 80, 320, 0.1122, 0.9944125, 0.6179),
    ("Aluminium", 300, 0.055, 167, 2400, 0.12342, 2.39749304, 0.2164),
]
r += 1
wsI.cell(r, 1, "C. CBAM products (Q0 kt 2024, g, EU exports X kt, price P $/t, F = fc+fp, G = np+no tCO2e/t, ER at $100)").font = BOLD
r += 1
for j, h in enumerate(["Product", "Q0", "g", "X", "P", "F", "G", "ER100"]):
    wsI.cell(r, 1 + j, h).font = BOLD
r += 1
P0 = r
for p in PROD:
    for j, v in enumerate(p):
        c = wsI.cell(r, 1 + j, v)
        if j:
            c.fill = GREEN
    r += 1
P1 = r - 1
for j, n in enumerate(["PT_Name", "PT_Q0", "PT_Growth", "PT_X", "PT_P", "PT_FuelFac", "PT_ProcFac", "PT_ER"]):
    col = "ABCDEFGH"[j]
    name(n, "Inputs!$%s$%d:$%s$%d" % (col, P0, col, P1))
wsI.column_dimensions["B"].width = 26
wsI.column_dimensions["C"].width = 52

# --------------------------------------------------------------------------------------------------------------
# Block: 8 products x 3 bundle variants (1A = 3A block, 3B with OBR, 3C with fund)
# --------------------------------------------------------------------------------------------------------------
wsB = wb.create_sheet("Block")
title(wsB, "CBAM block, all eight products, per bundle variant (Yr = 2030, NOPHASE)", 34)
COLS = ["Product", "Q0", "g", "X", "P", "F", "G", "beta", "Qb", "k_f", "k_p", "m (rebate/t)", "dp", "Q", "E_f", "E_p",
        "x_f", "ER_f", "ER_p", "E_post", "emrq", "emrq_proc", "rev_f $m", "rev_p $m", "rebate $m", "EI", "d", "obl",
        "E_base", "E_f base", "E_p base", "X*obl", "X*PEU*(F+G)", "covered"]
VARIANTS = [("1A", 1, 0, 0), ("3B", 1, 1, 0), ("3C", 1, 0, 1)]  # code, process flag, theta, phi
r = 3
BLK = {}
for code, flag, theta, phi in VARIANTS:
    wsB.cell(r, 1, "Bundle %s (flag=%d, theta=%d, phi=%d); 3A uses the 1A block" % (code, flag, theta, phi)).font = BOLD
    for c in range(1, len(COLS) + 1):
        wsB.cell(r, c).fill = BAND
    r += 1
    wsB.cell(r, 1, "flag")
    wsB.cell(r, 2, flag).fill = GREEN
    wsB.cell(r, 3, "theta")
    wsB.cell(r, 4, theta).fill = GREEN
    wsB.cell(r, 5, "phi")
    wsB.cell(r, 6, phi).fill = GREEN
    wsB.cell(r, 7, "tau_p")
    wsB.cell(r, 8, "=Tau*B%d" % r)
    wsB.cell(r, 9, "OBR rate")
    wsB.cell(r, 10, "=D%d*Tau" % r)
    wsB.cell(r, 11, "sigma_eff")
    wsB.cell(r, 12, "=F%d*SigmaEff" % r)
    pr = r
    TP, OBR, SG = "$H$%d" % pr, "$J$%d" % pr, "$L$%d" % pr
    r += 1
    for j, h in enumerate(COLS):
        wsB.cell(r, 1 + j, h).font = BOLD
    r += 1
    first = r
    for i in range(8):
        k = i + 1
        f = {
            "A": "=INDEX(PT_Name,%d)" % k, "B": "=INDEX(PT_Q0,%d)" % k, "C": "=INDEX(PT_Growth,%d)" % k,
            "D": "=INDEX(PT_X,%d)" % k, "E": "=INDEX(PT_P,%d)" % k, "F": "=INDEX(PT_FuelFac,%d)" % k,
            "G": "=INDEX(PT_ProcFac,%d)" % k, "H": "=-LN(1-INDEX(PT_ER,%d))/PStar" % k,
            "I": "=B{r}*(1+C{r})^(Yr-BaseYr)",
            "J": "=Tau*F{r}",
            "K": "=%s*G{r}" % TP,
            "L": "=MIN(%s,Tau)*F{r}+MIN(%s,%s)*G{r}" % (OBR, OBR, TP),
            "M": "=(J{r}+K{r}-L{r})/E{r}",
            "N": "=I{r}*(1+M{r})^EpsQ",
            "O": "=N{r}*F{r}/1000", "P": "=N{r}*G{r}/1000",
            "Q": "=EXP(Bf*(Tau+%s))" % SG,
            "R": "=O{r}*(Q{r}-1)",
            "S": "=-P{r}*(1-EXP(-H{r}*(%s+%s)))" % (TP, SG),
            "T": "=O{r}+R{r}+P{r}+S{r}",
            "U": "=(O{r}+P{r})*(1-(1+M{r})^(-EpsQ))",
            "V": "=P{r}*(1-(1+M{r})^(-EpsQ))",
            "W": "=(O{r}+R{r})*Tau",
            "X": "=(P{r}*(G{r}>0)+S{r})*%s" % TP,
            "Y": "=L{r}*N{r}/1000",
            "Z": "=IF(N{r}>0,1000*T{r}/N{r},0)",
            "AA": "=IF(N{r}>0,1000*(W{r}+X{r})/N{r}-L{r},0)",
            "AB": "=MAX(0,CBF*PEU*Z{r}-Ssel*MAX(0,AA{r}))",
            "AC": "=I{r}*(F{r}+G{r})/1000", "AD": "=I{r}*F{r}/1000", "AE": "=I{r}*G{r}/1000",
            "AF": "=D{r}*AB{r}", "AG": "=D{r}*CBF*PEU*(F{r}+G{r})",
            "AH": "=O{r}+P{r}*(G{r}>0)*$B$%d" % pr,
        }
        for col, frm in f.items():
            wsB["%s%d" % (col, r)] = frm.replace("{r}", str(r))
        r += 1
    last = r - 1
    wsB.cell(r, 1, "Total").font = BOLD
    for col in ["I", "N", "O", "P", "R", "S", "T", "U", "V", "W", "X", "Y", "AC", "AD", "AE", "AF", "AG", "AH"]:
        wsB["%s%d" % (col, r)] = "=SUM(%s%d:%s%d)" % (col, first, col, last)
        wsB["%s%d" % (col, r)].font = BOLD
    BLK[code] = dict(tot=r, first=first, clinker=first + 3, pr=pr)
    r += 3
wsB.column_dimensions["A"].width = 18

# --------------------------------------------------------------------------------------------------------------
# Example: one item per level
# --------------------------------------------------------------------------------------------------------------
wsE = wb.create_sheet("Example", 0)
title(wsE, "Worked example: ad hoc Table 2 rebuild v0.3 (REBUILD, NOPHASE, KappaMode=SCALE, 2030)", 9)
wsE["A2"] = ("One example per calculation level. Value = live formula; 'v0.3 value' = AdHocCalculations_Rebuild_v0.3 "
             "result (check). Last column: where the main version (CPAT_Industry_Kernel_Egypt_v0.15) differs; blank = same.")
wsE["A2"].alignment = WRAP
wsE.merge_cells("A2:I2")
wsE.row_dimensions[2].height = 30
HDR = ["Level", "Step", "Symbol", "Formula", "Value", "Unit", "v0.3 value", "Diff", "Main version (kernel v0.15)"]
for j, h in enumerate(HDR):
    c = wsE.cell(4, 1 + j, h)
    c.font = BOLD
    c.fill = BAND

c1, t1 = BLK["1A"]["clinker"], BLK["1A"]["tot"]
tB, tC = BLK["3B"]["tot"], BLK["3C"]["tot"]
B = "Block!"
# (level, step, symbol, formula text, excel formula, unit, check, main note); key -> row for references
STEPS = [
    ("0 Parameters", "Intensity share of fuel response", "s_int", "eps_F(1+eps_U)/(eps_U+eps_F(1+eps_U))",
     "=EpsF*(1+EpsU)/(EpsU+EpsF*(1+EpsU))", "", 1 / 3, "Not used: kernel keeps fuel intensity fixed"),
    ("0 Parameters", "Fuel-intensity semi-elasticity (EG1)", "b_f", "s_int*ln(IND1/IND0)/tau",
     "=E5*LN(IND1_EG1/IND0_EG1)/Tau", "per $/t", -0.0024956, "Not used (x_f = 1)"),
    ("0 Parameters", "Process semi-elasticity, clinker", "beta", "-ln(1-ER100)/P*, P* = 122",
     "=-LN(1-INDEX(PT_ER,4))/PStar", "per $/t", 0.0033465, "P* = 100 (USD2019, not deflated): beta = 0.00408"),
    ("1 CPAT inputs (EG1)", "Change in total GHG", "dGHG", "GHG1 - GHG0", "=GHG1_EG1-GHG0_EG1", "Mt", -38.877, ""),
    ("1 CPAT inputs (EG1)", "Change in IPPU", "dIPPU", "IPPU1 - IPPU0", "=IPPU1_EG1-IPPU0_EG1", "Mt", -11.955,
     "Replaced by kernel's own IPPU delta"),
    ("1 CPAT inputs (EG1)", "Change in industry energy CO2", "dInd", "IND1 - IND0", "=IND1_EG1-IND0_EG1", "Mt", -12.373,
     "Replaced by kernel's own industry delta"),
    ("2 Product (clinker, 1A)", "Baseline output 2030", "Qb", "Q0(1+g)^(Yr-2024)", "=%sI%d" % (B, c1), "kt", None, ""),
    ("2 Product (clinker, 1A)", "Carbon cost per t (fuel + process)", "k_f + k_p", "tau*F + tau_p*G",
     "=%sJ%d+%sK%d" % (B, c1, B, c1), "$/t", None, ""),
    ("2 Product (clinker, 1A)", "Relative price change", "dp", "(k_f+k_p-m)/P", "=%sM%d" % (B, c1), "", None, ""),
    ("2 Product (clinker, 1A)", "Output after price response", "Q", "Qb(1+dp)^eps_Q", "=%sN%d" % (B, c1), "kt", None, ""),
    ("2 Product (clinker, 1A)", "Fuel / process emissions at base intensity", "E_f + E_p", "Q(F+G)/1000",
     "=%sO%d+%sP%d" % (B, c1, B, c1), "Mt", None, ""),
    ("2 Product (clinker, 1A)", "Fuel-intensity response", "ER_f", "E_f(exp(b_f(tau+sigma))-1)", "=%sR%d" % (B, c1), "Mt",
     None, "0 (fuel intensity fixed)"),
    ("2 Product (clinker, 1A)", "Process-intensity response", "ER_p", "-E_p(1-exp(-beta(tau_p+sigma)))",
     "=%sS%d" % (B, c1), "Mt", None, "Same form, larger beta (P* = 100)"),
    ("2 Product (clinker, 1A)", "Output-channel process reduction", "emrq_proc", "E_p(1-(1+dp)^-eps_Q)",
     "=%sV%d" % (B, c1), "Mt", None, ""),
    ("2 Product (clinker, 1A)", "Domestic carbon revenue", "rev", "(E_f+ER_f)tau + (E_p+ER_p)tau_p",
     "=%sW%d+%sX%d" % (B, c1, B, c1), "$m", None, ""),
    ("2 Product (clinker, 1A)", "CBAM obligation per tonne", "obl", "max(0, CBF*P_EU*EI - S*max(0,d))",
     "=%sAB%d" % (B, c1), "$/t", None, "Headline uses FULL (CBF = 0.485 in 2030); NOPHASE shown in brackets"),
    ("3 Block (8 products, 1A)", "Block baseline emissions", "E_base", "sum Qb(F+G)/1000", "=%sAC%d" % (B, t1), "Mt",
     62.4236, ""),
    ("3 Block (8 products, 1A)", "Sum process-intensity response", "sum ER_p", "sum ER_p", "=%sS%d" % (B, t1), "Mt",
     -2.1254, ""),
    ("3 Block (8 products, 1A)", "Sum output-channel process reduction", "sum emrq_proc", "sum emrq_proc",
     "=%sV%d" % (B, t1), "Mt", -2.2985, ""),
    ("3 Block (8 products, 1A)", "[N] CBAM intensity change", "cbintch", "sum E_post / sum(E_f+E_p) - 1",
     "=%sT%d/(%sO%d+%sP%d)-1" % (B, t1, B, t1, B, t1), "", -0.0579, "-4.4% (no fuel-intensity channel)"),
    ("3 Block (8 products, 1A)", "[O] CBAM obligation per unit change", "cbobchu", "sum X*obl / sum X*CBF*P_EU(F+G) - 1",
     "=%sAF%d/%sAG%d-1" % (B, t1, B, t1), "", -0.2434, "-43.3% FULL (-22.8% NOPHASE)"),
    ("3 Block (8 products, 1A)", "[T] Block emissions change", "emrt/E_base", "(sum emrq + sum ER_f + sum ER_p)/E_base",
     "=(%sU%d+%sR%d+%sS%d)/%sAC%d" % (B, t1, B, t1, B, t1, B, t1), "", -0.1156, ""),
    ("4 National (1A)", "[J] Coverage", "J", "(ENR0 + flag*E_p base)/GHG0",
     "=(ENR0_EG1+%sAE%d)/GHG0_EG1" % (B, t1), "", 0.6320, ""),
    ("4 National (1A)", "[K] Total reduction", "K", "dGHG - dIPPU + IPPU_other(0) + sum ER_p + sum emrq_proc",
     "=E8-E9+0+E22+E23", "Mt", -31.3458,
     "Composed as dGHG_CPAT + (dIPPU_kernel - dIPPU_CPAT) + (dInd_kernel - dInd_CPAT): -29.8"),
    ("4 National (1A)", "[L] Reduction share", "L", "K / GHG0", "=E28/GHG0_EG1", "", -0.0528, ""),
    ("4 National (1A)", "[P] Net revenue", "P", "CPAT fuel receipts + sum rev_p/1000 - rebates",
     "=REVCTAX_EG1+%sX%d/1000-%sY%d/1000" % (B, t1, B, t1), "$bn", 6.8380, "6.87"),
    ("4 National (1A)", "[Q] Deaths avoided", "Q", "CPAT deaths (own scenario)", "=DEATHS_EG1", "#", 1564, "1,441"),
    ("4 National (1A)", "[AR] Delta net revenue", "AR", "CPAT dRev - (P_gross - P)", "=DREV_EG1-0", "$bn", 10.43, ""),
    ("5 Variant 3A (EG3)", "Priced share of industry energy CO2", "kappa", "min(1, effcpt/cpt * ENR0/IND0)",
     "=MIN(1,EFFCPT_EG3/CPT_EG3*ENR0_EG3/IND0_EG3)", "", 0.5374,
     "No kappa/SCALE: kernel prices only its four sectors (coverage 8.6%)"),
    ("5 Variant 3A (EG3)", "SCALE: adjusted industry delta", "dInd_adj", "dInd / kappa",
     "=(IND1_EG3-IND0_EG3)/E33", "Mt", None, "n/a"),
    ("5 Variant 3A (EG3)", "SCALE: adjusted IPPU delta", "dIPPU_adj", "dIPPU / kappa",
     "=(IPPU1_EG3-IPPU0_EG3)/E33", "Mt", -15.3748, "n/a"),
    ("5 Variant 3A (EG3)", "SCALE: adjusted GHG delta", "dGHG_adj", "dGHG + (1/kappa-1)(dInd+dIPPU)",
     "=(GHG1_EG3-GHG0_EG3)+(1/E33-1)*((IND1_EG3-IND0_EG3)+(IPPU1_EG3-IPPU0_EG3))", "Mt", -36.0164, "n/a"),
    ("5 Variant 3A (EG3)", "[J] Coverage (kappa = 1 under SCALE)", "J", "(IND0 + E_p base)/GHG0",
     "=(IND0_EG3+%sAE%d)/GHG0_EG3" % (B, t1), "", 0.2084, "8.6%"),
    ("5 Variant 3A (EG3)", "[K] Total reduction", "K", "dGHG_adj - dIPPU_adj + sum ER_p + sum emrq_proc (1A block)",
     "=E36-E35+E22+E23", "Mt", -25.0655, "-11.8"),
    ("5 Variant 3A (EG3)", "[Q] Deaths avoided", "Q", "CPAT deaths * dENR_adj/dENR",
     "=DEATHS_EG3*((ENR1_EG3-ENR0_EG3)+(1/E33-1)*(IND1_EG3-IND0_EG3))/(ENR1_EG3-ENR0_EG3)", "#", 850.3418, "278"),
    ("5 Variant 3B (OBR)", "OBR correction", "D_obr", "-(1-s_int)*w*dInd_adj, w = E_f base/IND0",
     "=-(1-E5)*MIN(1,%sAD%d/IND0_EG3)*E34" % (B, tB), "Mt", 3.2911,
     "Rebate capped by sector shares; no D_obr term"),
    ("5 Variant 3B (OBR)", "[K] Total reduction", "K", "dGHG_adj - dIPPU_adj + sum ER_p + sum emrq_proc + D_obr",
     "=E36-E35+%sS%d+%sV%d+E40" % (B, tB, B, tB), "Mt", -19.6263, "-8.6"),
    ("5 Variant 3B (OBR)", "[P] Net revenue", "P", "REVCTAX/kappa + sum rev_p/1000 - rebates",
     "=REVCTAX_EG3/E33+%sX%d/1000-%sY%d/1000" % (B, tB, B, tB), "$bn", 1.1331, "-0.02"),
    ("5 Variant 3C (fund)", "Fund fuel abatement outside block", "F_fund", "IND1_adj(exp(b_f*sigma)-1)",
     "=(IND0_EG3+E34)*(EXP(E6*SigmaEff)-1)", "Mt", -3.5571,
     "Fund solved as shadow price (~150 $/t) spending whole fund on block"),
    ("5 Variant 3C (fund)", "[K] Total reduction", "K", "dGHG_adj - dIPPU_adj + sum ER_p + sum emrq_proc + F_fund",
     "=E36-E35+%sS%d+%sV%d+E43" % (B, tC, B, tC), "Mt", -30.6038, "-23.9"),
]
r = 5
for lvl, step, sym, ftxt, frm, unit, chk, main in STEPS:
    wsE.cell(r, 1, lvl)
    wsE.cell(r, 2, step)
    wsE.cell(r, 3, sym)
    wsE.cell(r, 4, ftxt)
    wsE.cell(r, 5, frm).number_format = "0.0000"
    wsE.cell(r, 6, unit)
    if chk is not None:
        wsE.cell(r, 7, chk).number_format = "0.0000"
        wsE.cell(r, 8, "=E%d-G%d" % (r, r)).number_format = "0.0000"
    m = wsE.cell(r, 9, main or "")
    if main:
        m.fill = YELLOW
    r += 1
name("Bf", "Example!$E$6")
CHK_ROW = r + 1
wsE.cell(CHK_ROW, 2, "Max |Diff| (check values rounded to 4 dp)").font = BOLD
wsE.cell(CHK_ROW, 8, "=MAX(INDEX(ABS(H5:H%d),0))" % (r - 1)).number_format = "0.0000"
wsE.cell(CHK_ROW, 9, '=IF(H%d<0.001,"PASS","CHECK")' % CHK_ROW)
for col, w in zip("ABCDEFGHI", [24, 40, 14, 48, 12, 8, 12, 10, 60]):
    wsE.column_dimensions[col].width = w
wsE.freeze_panes = "A5"

wb.save(OUT)
print("saved", OUT, "check row", CHK_ROW)
