"""Build the seed CPAT country data workbook (cpat_country_data.xlsx).

One long table, one row per country, year and variable, so that Python and
Excel read country data from a single shared file instead of each workbook
holding it. The file is uploaded to the CPAT SharePoint folder `CountryData`
(see README.md); after that it is edited there, not rebuilt from here.

Seed content (mitigation module):
  - domestic prices: all countries, 2019-2024, from Prices_dom in
    cpat_excel_new/working_version/CPAT-AI-Mitigation-MVP-v1.00.xlsx
    (includes the user-corrected Egypt block, flagged in the source column);
  - energy balances: Egypt 2022 from the legacy Balances sheet
    (cpat_excel_original, the NoPropData workbook holds Egypt only);
  - base-year energy use by subsector and fuel: Egypt 2022 from EnergyCons
    in the MVP workbook.

Run from the repo root:  python country_data/build_country_data.py
Needs pandas, openpyxl, pyxlsb (as cpat_coded/requirements.txt).
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pyxlsb import open_workbook

ROOT = Path(__file__).resolve().parents[1]
MVP = ROOT / 'cpat_excel_new' / 'working_version' / 'CPAT-AI-Mitigation-MVP-v1.00.xlsx'
LEGACY = ROOT / 'cpat_excel_original' / 'CPAT 1.0pre_456_NoPropData.xlsb'
OUT = Path(__file__).resolve().parent / 'cpat_country_data.xlsx'

VINTAGE = '2026-10-10'
CHANGED_FILL = '00FFFF00'  # bright yellow = cell changed by the corrected Egypt block

SRC_PRICES = 'CPAT-AI-Mitigation-MVP v1.00 Prices_dom (IMF price dataset via legacy CPAT)'
SRC_PRICES_EGY = 'Corrected Egypt price block (user, 2026-10-09) via MVP v1.00 Prices_dom'
SRC_BAL = 'Legacy CPAT 1.0pre_456 Balances sheet'
SRC_ENER = 'MVP v1.00 EnergyCons = legacy Mitigation rows 583-602 (derived from Balances)'

COLUMNS = ['iso3', 'year', 'variable', 'value', 'unit', 'source', 'vintage']

# NORMS section 2
HEADER_FILL = PatternFill('solid', fgColor='EBF1DE')
BAND_FILL = PatternFill('solid', fgColor='00B050')
FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)
TITLE = Font(name='Arial', size=14, bold=True, color='FFFFFF')


def read_prices():
    """Prices_dom (wide: country_year x code) -> long rows and variable labels."""
    ws = openpyxl.load_workbook(MVP)['Prices_dom']
    ncol = ws.max_column
    labels, rows = {}, []
    fuel = sector = None
    for c in range(6, ncol + 1):
        if ws.cell(1, c).value:
            fuel, sector = ws.cell(1, c).value, None
        if ws.cell(2, c).value:
            sector = ws.cell(2, c).value
        indicator = str(ws.cell(3, c).value or '').replace('\n', ' ').strip()
        code = ws.cell(5, c).value
        labels[code] = (' / '.join(x for x in (fuel, sector, indicator) if x), ws.cell(4, c).value or '')
    for r in range(6, ws.max_row + 1):
        iso3, year = ws.cell(r, 3).value, ws.cell(r, 5).value
        for c in range(6, ncol + 1):
            cell = ws.cell(r, c)
            if cell.value is None:
                continue
            code = ws.cell(5, c).value
            src = SRC_PRICES_EGY if cell.fill.fgColor.rgb == CHANGED_FILL else SRC_PRICES
            rows.append([iso3, int(year), code, float(cell.value), labels[code][1], src, VINTAGE])
    variables = [[code, 'prices_dom', desc, unit] for code, (desc, unit) in labels.items()]
    return rows, variables


def read_balances():
    """Legacy Balances (country, year, measure, flow x fuel columns) -> long rows."""
    rows, header = [], None
    with open_workbook(str(LEGACY)) as wb, wb.get_sheet('Balances') as sh:
        for row in sh.rows():
            vals = [c.v for c in row]
            if vals[0] == 'country_code':
                header = vals
                continue
            if header is None or not vals[0]:
                continue
            iso3, year, unit, flow = vals[0], int(vals[2]), vals[3], vals[4]
            for fuel, v in zip(header[5:], vals[5:]):
                if v is not None:
                    rows.append([iso3, year, f'bal.{flow}.{fuel}', float(v), unit, SRC_BAL, VINTAGE])
    seen = {}
    for r in rows:
        _, flow, fuel = r[2].split('.')
        seen[r[2]] = [r[2], 'balances', f'Energy balance: flow {flow}, fuel {fuel}', r[4]]
    return rows, list(seen.values())


def read_energy_use():
    """MVP EnergyCons (key, country, year, subsector, fuel, ktoe) -> long rows."""
    ws = openpyxl.load_workbook(MVP, read_only=True)['EnergyCons']
    rows, variables = [], {}
    for key, iso3, year, sub, fuel, ktoe, _ in ws.iter_rows(min_row=4, values_only=True):
        if not key:
            continue
        code = f'mit.ener.{sub}.{fuel}'
        rows.append([iso3, int(year), code, float(ktoe), 'ktoe', SRC_ENER, VINTAGE])
        variables[code] = [code, 'energy_use', f'Base-year energy use: subsector {sub}, fuel {fuel}', 'ktoe']
    return rows, list(variables.values())


def write_sheet(ws, header, rows, table_name, widths):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ref = f'A1:{openpyxl.utils.get_column_letter(len(header))}{len(rows) + 1}'
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(table)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, BOLD
    ws.freeze_panes = 'A2'


def write_readme(ws, n_rows, n_countries):
    lines = [
        ('CPAT country data', None),
        (None, None),
        ('Purpose', 'Shared country data for CPAT (Python and Excel). Models look values up here instead of holding them. '
                    'Lives on the CPAT SharePoint drive, folder CountryData; do not rename or move the file (every user and workbook points at this path).'),
        ('Table', f'Sheet Data, Excel table CountryData: one row per iso3 + year + variable ({n_rows:,} rows, {n_countries} countries in this seed).'),
        ('Columns', 'iso3 = ISO3 country code; year = data year; variable = CPAT code (see Variables); value = number; '
                    'unit = unit of value; source = where the number comes from; vintage = date the row was added or last changed (YYYY-MM-DD).'),
        ('Lookup key', 'iso3|year|variable, e.g. EGY|2022|mit.sp.nga.ind. Each key must occur once (the Python loader refuses duplicates).'),
        ('Variable groups', 'prices_dom = domestic prices, taxes and coefficients (legacy Prices_dom codes, nominal units of the data year); '
                            'balances = energy balances bal.<flow>.<fuel> (legacy Balances codes); '
                            'energy_use = base-year energy use mit.ener.<subsector>.<fuel> in ktoe (as the Mitigation module uses it).'),
        ('Adding or changing data', 'Add rows at the bottom of the table (it grows automatically). Fill every column. '
                                    'To correct a value, overwrite it and update source and vintage; SharePoint version history keeps the old file. '
                                    'New variable? Add it to Variables first.'),
        ('Do not', 'Insert columns, rename the table or sheets, merge cells, or type formulas in Data (readers see stored values only).'),
        ('Licensed data', 'Balances derived from IEA data are licensed: keep this file on WB SharePoint only, with access limited to the CPAT team.'),
        ('How to use', 'See country_data/README.md in the CPAT-AI-Local repo (Python loader and Excel Power Query steps).'),
        (None, None),
        ('Change log', None),
        ('Date', 'Change'),
        (VINTAGE, 'Seed: domestic prices all countries 2019-2024 (MVP v1.00 Prices_dom, corrected Egypt block flagged in source), '
                  'Egypt 2022 energy balances (legacy Balances) and Egypt 2022 base-year energy use (MVP EnergyCons). '
                  'Built by country_data/build_country_data.py.'),
    ]
    for label, text in lines:
        ws.append([label, text])
    ws['A1'].font = TITLE
    for c in ('A1', 'B1'):
        ws[c].fill = BAND_FILL
    for row in ws.iter_rows(min_row=2):
        row[0].font = BOLD
        row[1].font = FONT
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        row[0].alignment = Alignment(vertical='top')
    for c in ws[14]:
        c.fill = HEADER_FILL
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 110


def main():
    p_rows, p_vars = read_prices()
    b_rows, b_vars = read_balances()
    e_rows, e_vars = read_energy_use()
    rows = p_rows + b_rows + e_rows
    keys = [(r[0], r[1], r[2]) for r in rows]
    assert len(keys) == len(set(keys)), 'duplicate iso3|year|variable in seed'

    wb = openpyxl.Workbook()
    write_readme(wb.active, len(rows), len({r[0] for r in rows}))
    wb.active.title = 'ReadMe'
    write_sheet(wb.create_sheet('Data'), COLUMNS, rows, 'CountryData', [8, 7, 24, 14, 16, 60, 12])
    write_sheet(wb.create_sheet('Variables'), ['variable', 'group', 'description', 'unit'],
                p_vars + b_vars + e_vars, 'Variables', [24, 12, 60, 16])
    wb.save(OUT)
    print(f'{OUT.name}: {len(p_rows)} price rows, {len(b_rows)} balance rows, {len(e_rows)} energy-use rows; '
          f'{sum(r[5] == SRC_PRICES_EGY for r in p_rows)} corrected Egypt price cells')


if __name__ == '__main__':
    main()
