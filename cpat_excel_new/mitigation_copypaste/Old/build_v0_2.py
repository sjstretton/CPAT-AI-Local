"""Build CPAT_Mitigation_CopyPaste_v0.2.xlsx (design 2: time across, scenarios as column groups).

Price -> fuel-use prototype of the CPAT mitigation module with fully copy-pasteable formulas.

v0.2 structure:
- Data input is a separate step: sheet `Inputs` has one row per fuel|subsector (same taxonomy as
  Mitigation columns B:C) holding every lookup (mappings, elasticities for the selected income group,
  base prices, base tax, EF, base-year fuel use). `Mitigation` reads it with one short INDEX/MATCH in
  four hidden parameter columns D:G and in the base-year column; calculation cells hold no lookups.
- Parameters are global (not scenario-specific); scenarios differ only in their macro rows.
- 2023-2034 hold plain formulas; the last year (2035) calls named LAMBDAs (PRETAX, TAX, POSTTAX,
  FUELUSE) in a distinct colour, so they can be dragged back over the row if preferred.
- Base year 2022 sits in column L, as in the legacy Mitigation sheet.

Reads data/*.csv (made by extract_data_v0_1.py). Writes formulas only (no cached values).

    python build_v0_2.py
"""
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

VERSION = '0.2'
HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'data')
OUT = os.path.join(HERE, f'CPAT_Mitigation_CopyPaste_v{VERSION}.xlsx')

# ---------------------------------------------------------------- dimensions
SUBSECTORS = [  # code, name, group, elasticity sector, EF sector, coal/gas price sector, carbon-price coverage
    ('rod', 'Road', 'tra', 'tra', 'rod', 'res', 1),
    ('ral', 'Rail', 'tra', 'tra', 'rod', 'ind', 1),
    ('avi', 'Domestic aviation', 'tra', 'tra', 'rod', 'ind', 1),
    ('nav', 'Domestic navigation', 'tra', 'tra', 'rod', 'ind', 1),
    ('res', 'Residential', 'bld', 'res', 'res', 'res', 1),
    ('foo', 'Food & forestry', 'bld', 'ind', 'res', 'ind', 1),
    ('srv', 'Services (public & private)', 'bld', 'srv', 'res', 'ind', 1),
    ('mch', 'Mining & chemicals', 'ind', 'ind', 'ind', 'ind', 1),
    ('irn', 'Iron & steel', 'ind', 'ind', 'ind', 'ind', 1),
    ('nfm', 'Other metals', 'ind', 'ind', 'ind', 'ind', 1),
    ('mac', 'Machinery', 'ind', 'ind', 'ind', 'ind', 1),
    ('cem', 'Cement', 'ind', 'ind', 'ind', 'ind', 1),
    ('omn', 'Other manufacturing', 'ind', 'ind', 'ind', 'ind', 1),
    ('cst', 'Construction', 'ind', 'ind', 'ind', 'ind', 1),
    ('ftr', 'Fuel transformation', 'ind', 'ind', 'ind', 'ind', 1),
    ('oen', 'Other energy use', 'oen', 'ind', 'ind', 'ind', 0),
]
FUELS = [  # code, name, elasticity fuel, AEEI fuel, price-sector rule, price unit
    ('coa', 'Coal', 'coa', 'coa', 'sub', '$/GJ'),
    ('nga', 'Natural gas', 'nga', 'nga', 'sub', '$/GJ'),
    ('gso', 'Gasoline', 'gso', 'oil', 'all', '$/liter'),
    ('die', 'Diesel', 'die', 'oil', 'all', '$/liter'),
    ('lpg', 'LPG', 'oil', 'oil', 'all', '$/liter'),
    ('ker', 'Kerosene', 'oil', 'oil', 'all', '$/liter'),
    ('oop', 'Other oil products', 'oil', 'oil', 'all', '$/bbl'),
    ('bio', 'Biomass', 'bio', 'bio', 'all', '$/GJ'),
]
NS, NF = len(SUBSECTORS), len(FUELS)
NROW = NS * NF                          # 128 rows per block (subsector outer, fuel inner)
SNAMES = {s[0]: s[1] for s in SUBSECTORS}
FNAMES = {f[0]: f[1] for f in FUELS}
PAIRS = [(s[0], f[0]) for s in SUBSECTORS for f in FUELS]

BASE_YEAR, LAST_YEAR = 2022, 2035
YEARS = list(range(BASE_YEAR + 1, LAST_YEAR + 1))
SCENARIOS = [(1, 'Baseline (no carbon price)', {}),
             (2, 'Carbon price $20/tCO2 from 2027', {y: 20 for y in YEARS if y >= 2027})]

# ---------------------------------------------------------------- Inputs sheet layout (one row per fuel|subsector)
IN_R0 = 5                               # first data row
IN_R1 = IN_R0 + NROW - 1
INPUT_COLS = [  # column header, short code
    ('Key (fuel|subsector)', 'key'), ('Fuel', 'fuel'), ('Subsector', 'sub'),
    ('Elasticity sector', 'esec'), ('Elasticity fuel', 'efuel'), ('AEEI fuel', 'afuel'),
    ('Price code', 'pcode'), ('EF sector', 'efsec'), ('GJ per price unit', 'gj'), ('Coverage', 'cov'),
    ('eps_Y GDP', 'eY'), ('eps_U usage', 'eU'), ('eps_F efficiency', 'eF'), ('alpha AEEI', 'a'),
    ('Pre-tax price growth', 'grow'), ('Base pre-tax price ($/GJ)', 'p0'), ('Base tax ($/GJ)', 't0'),
    ('EF (tCO2/GJ)', 'ef'), ('Base-year fuel use (ktoe)', 'f0'),
]
IC = {code: L(j) for j, (_, code) in enumerate(INPUT_COLS, 1)}

# ---------------------------------------------------------------- Mitigation layout
PARAM_COLS = ['D', 'E', 'F', 'G']       # hidden by default
DESC = {'H': 'Description', 'I': 'Unit', 'J': 'Source', 'K': 'Output Code'}
COL_G0 = 12                             # L: base year of scenario group 1 (= legacy Mitigation column L, 2022)
NY = 1 + len(YEARS)                     # base year + projection years
GW = NY + 1                             # group width incl. one gap column
R_YEAR, R_NAME, R_GDP, R_CP = 4, 5, 6, 7
BLOCK_H = NROW + 2                      # band row + 128 rows + blank row
B_PRE, B_TAX, B_POST, B_USE = 9, 9 + BLOCK_H, 9 + 2 * BLOCK_H, 9 + 3 * BLOCK_H
B_TOT = 9 + 4 * BLOCK_H

BLOCKS = {  # band row, band text, code, unit, short name, parameters in D:G
    'pre': (B_PRE, 'Pre-tax price: p(t) = p(t-1) x (1 + growth)', 'pbt', '$/GJ', 'Pre-tax price', ['grow']),
    'tax': (B_TAX, 'Tax: base tax + carbon price x EF x coverage', 'tax', '$/GJ', 'Tax', ['t0', 'ef', 'cov']),
    'post': (B_POST, 'Post-tax price = pre-tax price + tax', 'ptp', '$/GJ', 'Post-tax price', []),
    'use': (B_USE, 'Fuel use (CPAT documentation 3.3.3)', 'enc', 'ktoe', 'Fuel use', ['eY', 'eU', 'eF', 'a']),
}
PARAM_LABELS = {'grow': 'growth', 't0': 'base tax', 'ef': 'EF', 'cov': 'coverage',
                'eY': 'eps_Y', 'eU': 'eps_U', 'eF': 'eps_F', 'a': 'alpha'}

# ---------------------------------------------------------------- named LAMBDAs (used in the last year only)
LAMBDAS = {
    'PRETAX': (['prev', 'growth'], 'prev*(1+growth)'),
    'TAX': (['base_tax', 'cp', 'ef', 'cov'], 'base_tax+cp*ef*cov'),
    'POSTTAX': (['pre_tax_p', 'tax_p'], 'pre_tax_p+tax_p'),
    'FUELUSE': (['f_prev', 'p_now', 'p_prev', 'gdp_g', 'eps_y', 'eps_u', 'eps_f', 'alpha'],
                'f_prev*(1/(1+alpha))^(1+eps_u)*(1+gdp_g)^eps_y*(p_now/p_prev)^eps_u'
                '*(p_now/p_prev)^(eps_f*(1+eps_u))'),
}
LAMBDA_NOTES = {
    'PRETAX': 'Pre-tax price: previous year x (1 + growth)',
    'TAX': 'Tax ($/GJ): base tax + carbon price ($/tCO2) x EF (tCO2/GJ) x coverage',
    'POSTTAX': 'Post-tax price: pre-tax price + tax',
    'FUELUSE': 'Fuel use, CPAT documentation 3.3.3 (no Covid factor, no shadow price)',
}


def lambda_xml(name):
    """Excel file encoding of a LAMBDA: _xlfn.LAMBDA(_xlpm.x, ..., body with _xlpm.x)."""
    params, body = LAMBDAS[name]
    for prm in sorted(params, key=len, reverse=True):
        body = re.sub(rf'(?<![A-Za-z_.]){prm}(?![A-Za-z_\d(])', f'_xlpm.{prm}', body)
    return '_xlfn.LAMBDA(' + ','.join(f'_xlpm.{prm}' for prm in params) + ',' + body + ')'


# ---------------------------------------------------------------- styles (NORMS section 2)
def fill(hex_):
    return PatternFill('solid', start_color=hex_, end_color=hex_)


F_TITLE, F_BAND, F_INPUT, F_CODE = fill('00B050'), fill('92D050'), fill('EBF1DE'), fill('DCE6F1')
F_BASE, F_UNUSED, F_CHECK, F_LAMBDA = fill('DDD9C4'), fill('F2F2F2'), fill('FFF2CC'), fill('FCE4D6')
FONT = Font(name='Arial', size=9)
FONT_B = Font(name='Arial', size=9, bold=True)
FONT_T = Font(name='Arial', size=14, bold=True, color='FFFFFF')
FONT_LAMBDA = Font(name='Arial', size=9, color='9C0006')


def put(ws, cell, value, font=FONT, fill_=None, fmt=None):
    c = ws[cell]
    c.value = value
    c.font = font
    if fill_:
        c.fill = fill_
    if fmt:
        c.number_format = fmt
    return c


def title(ws, text, width):
    for j in range(1, width + 1):
        ws.cell(2, j).fill = F_TITLE
    put(ws, 'A2', text, FONT_T, F_TITLE)


def band(ws, row, text, width):
    for j in range(1, width + 1):
        ws.cell(row, j).fill = F_BAND
    put(ws, f'A{row}', text, FONT_B, F_BAND)


def read_csv(name):
    with open(os.path.join(DATA, name), encoding='utf-8') as f:
        return list(csv.reader(f))


def num(v):
    if v in ('', None):
        return None
    try:
        return float(v)
    except ValueError:
        return v


# ---------------------------------------------------------------- data sheets (unchanged from v0.1)
def data_sheet(wb, name, title_text, header, rows, source):
    ws = wb.create_sheet(name)
    title(ws, title_text, len(header))
    put(ws, 'A1', f'Source: {source}', FONT)
    for j, h in enumerate(header, 1):
        put(ws, f'{L(j)}3', num(h) if isinstance(h, str) and h.isdigit() else h, FONT_B, F_INPUT)
    for i, r in enumerate(rows, 4):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, num(v)).font = FONT
    ws.freeze_panes = 'B4'
    ws.column_dimensions['A'].width = 18


def build_data_sheets(wb):
    c = read_csv('countries.csv')
    data_sheet(wb, 'Countries', 'Countries: code, region and income group', c[0], c[1:],
               'legacy Elasticities rows 138-141 (country columns)')
    e = read_csv('elasticities.csv')
    data_sheet(wb, 'Elasticities', 'Elasticities by income group (legacy CPAT values, CPAT sector and fuel groups)',
               e[0], e[1:], "legacy Elasticities rows 143-282 ('Simple' option, income-group specific)")
    hdr, rows = read_csv('prices_dom_header.csv'), read_csv('prices_dom.csv')
    ws = wb.create_sheet('Prices_dom')  # legacy Prices_dom layout (NORMS 1.3): bands rows 1-4, codes row 5
    for i, r in enumerate(hdr[1:5], 1):
        for j, v in enumerate(r, 1):
            if v:
                ws.cell(i, j, v).font = FONT
    for j, h in enumerate(rows[0], 1):
        put(ws, f'{L(j)}5', h, FONT_B, F_INPUT)
    for i, r in enumerate(rows[1:], 6):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, num(v) if j > 4 else (int(float(v)) if j == 5 else v)).font = FONT
    ws.freeze_panes = 'F6'
    u = read_csv('energy_use.csv')
    data_sheet(wb, 'EnergyCons', 'Base-year energy use by subsector and fuel (ktoe)', u[0], u[1:],
               'legacy Mitigation rows 583-602 via kernel v1.6 Data_Energy (Egypt only so far)')
    w = read_csv('weo.csv')
    data_sheet(wb, 'WEO', 'Macro series by country (real GDP growth)', w[0], w[1:],
               'legacy Mitigation row 2446 via kernel v1.6 Data_Macro (Egypt only so far)')
    f = read_csv('ef_co2.csv')
    data_sheet(wb, 'EF_GHG', 'Fuel CO2 emission factors, IIASA (tCO2/GJ), by country, fuel and EF sector',
               f[0], f[1:], 'legacy Mitigation rows 980-1009 (Egypt only so far; no inventory adjustment)')


# ---------------------------------------------------------------- Mapping, Settings, Inputs
def build_mapping(wb):
    ws = wb.create_sheet('Mapping')
    title(ws, 'Mapping: subsectors and fuels to CPAT parameter groups', 8)
    put(ws, 'A3', 'Subsectors', FONT_B)
    for j, h in enumerate(['Code', 'Name', 'Group', 'Elasticity sector', 'EF sector',
                           'Coal/gas price sector', 'Carbon-price coverage'], 1):
        put(ws, f'{L(j)}4', h, FONT_B, F_INPUT)
    for i, s in enumerate(SUBSECTORS, 5):
        for j, v in enumerate(s, 1):
            put(ws, f'{L(j)}{i}', v, FONT, F_INPUT if j > 2 else None)
    units = {r[0]: float(r[1]) for r in read_csv('fuel_units.csv')[1:]}
    put(ws, 'A22', 'Fuels', FONT_B)
    for j, h in enumerate(['Code', 'Name', 'Elasticity fuel', 'AEEI fuel', 'Price sector rule',
                           'Price unit', 'GJ per price unit'], 1):
        put(ws, f'{L(j)}23', h, FONT_B, F_INPUT)
    for i, f in enumerate(FUELS, 24):
        for j, v in enumerate(list(f) + [units[f[0]]], 1):
            put(ws, f'{L(j)}{i}', v, FONT, F_INPUT if j > 2 else None)
    put(ws, 'A33', "Price sector rule: 'sub' = use the subsector's coal/gas price sector (res or ind, "
                   "as CPAT: rod and res use residential coal/gas prices); 'all' = the fuel's single price.", FONT)
    put(ws, 'A34', 'GJ per price unit: legacy Mitigation rows 980-1009, conversion factor (volume unit to GJ).', FONT)
    put(ws, 'A35', 'Coverage: CPAT documentation 3.3.4.6 - no carbon taxes in other energy use (oen).', FONT)
    ws.column_dimensions['B'].width = 26
    for col in 'CDEFG':
        ws.column_dimensions[col].width = 16


def build_settings(wb):
    ws = wb.create_sheet('Settings')
    title(ws, 'Settings', 6)
    rows = [(4, 'Country code', 'EGY', True),
            (5, 'Country name', '=INDEX(Countries!$B$4:$B$1000,MATCH($C$4,Countries!$A$4:$A$1000,0))', False),
            (6, 'Income group (selects elasticities)',
             '=INDEX(Countries!$D$4:$D$1000,MATCH($C$4,Countries!$A$4:$A$1000,0))', False),
            (7, 'Base year', BASE_YEAR, True)]
    for r, lab, val, inp in rows:
        put(ws, f'B{r}', lab)
        put(ws, f'C{r}', val, FONT, F_INPUT if inp else None)
    put(ws, 'B10', 'Version log', FONT_B)
    for j, h in enumerate(['Version', 'Date', 'Description', 'Max abs regression diff'], 2):
        put(ws, f'{L(j)}11', h, FONT_B, F_INPUT)
    log = [('0.1', '2026-10-08', 'First build: design 2 price -> fuel-use prototype, Egypt base data, 2 scenarios.',
            'n/a (first version)'),
           (VERSION, '2026-10-08', 'Inputs sheet by fuel|subsector (lookups separated from formulas); 4 hidden '
                                   'parameter columns D:G, global parameters; base year 2022 in column L; '
                                   'named LAMBDAs in 2035.', '0 (all Mitigation results vs v0.1)')]
    for i, row in enumerate(log, 12):
        for j, v in enumerate(row, 2):
            put(ws, f'{L(j)}{i}', v)
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 90
    ws.column_dimensions['E'].width = 30


def build_inputs(wb):
    ws = wb.create_sheet('Inputs')
    title(ws, 'Inputs by fuel | subsector: every lookup used by Mitigation (data step, no calculations)',
          len(INPUT_COLS))
    put(ws, 'A3', 'One row per fuel | subsector (same taxonomy as Mitigation columns B:C). Mitigation reads these '
                  'columns with INDEX/MATCH on the key. Green = input.', FONT)
    for j, (h, _) in enumerate(INPUT_COLS, 1):
        c = put(ws, f'{L(j)}4', h, FONT_B, F_INPUT)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    sub_ = lambda col, r: f'INDEX(Mapping!${col}$5:${col}$20,MATCH($C{r},Mapping!$A$5:$A$20,0))'
    fuel_ = lambda col, r: f'INDEX(Mapping!${col}$24:${col}$31,MATCH($B{r},Mapping!$A$24:$A$31,0))'

    def elast(typ, fcol, r):
        return (f'=INDEX(Elasticities!$F$4:$I$200,MATCH("{typ}|"&${fcol}{r}&"|"&${IC["esec"]}{r},'
                f'Elasticities!$A$4:$A$200,0),MATCH(Settings!$C$6,Elasticities!$F$3:$I$3,0))')

    def price(prefix, r):
        return (f'INDEX(Prices_dom!$A$6:$CU$2000,MATCH(Settings!$C$4&Settings!$C$7,Prices_dom!$A$6:$A$2000,0),'
                f'MATCH("{prefix}."&${IC["pcode"]}{r},Prices_dom!$A$5:$CU$5,0))')

    for i, (s, f) in enumerate(PAIRS):
        r = IN_R0 + i
        vals = {
            'key': f'=$B{r}&"|"&$C{r}', 'fuel': f, 'sub': s,
            'esec': '=' + sub_('D', r), 'efuel': '=' + fuel_('C', r), 'afuel': '=' + fuel_('D', r),
            'pcode': f'=$B{r}&"."&IF({fuel_("E", r)}="sub",{sub_("F", r)},"all")',
            'efsec': '=' + sub_('E', r), 'gj': '=' + fuel_('G', r), 'cov': '=' + sub_('G', r),
            'eY': elast('inc', IC['efuel'], r), 'eU': elast('usg', IC['efuel'], r),
            'eF': elast('eff', IC['efuel'], r), 'a': elast('aei', IC['afuel'], r),
            'grow': 0,
            'p0': f'={price("mit.sp", r)}/${IC["gj"]}{r}',
            't0': f'=({price("mit.rp", r)}-{price("mit.sp", r)})/${IC["gj"]}{r}',
            'ef': (f'=INDEX(EF_GHG!$E$4:$E$1000,MATCH(Settings!$C$4&"|"&$B{r}&"|"&${IC["efsec"]}{r},'
                   f'EF_GHG!$A$4:$A$1000,0))'),
            'f0': (f'=INDEX(EnergyCons!$F$4:$F$5000,MATCH(Settings!$C$4&"|"&Settings!$C$7&"|"&$C{r}&"|"&$B{r},'
                   f'EnergyCons!$A$4:$A$5000,0))'),
        }
        for _, code in INPUT_COLS:
            fmt = {'p0': '0.000', 't0': '0.000', 'ef': '0.0000', 'f0': '#,##0.0', 'grow': '0.0%'}.get(code)
            if code in ('eY', 'eU', 'eF'):
                fmt = '0.00'
            if code == 'a':
                fmt = '0.0%'
            put(ws, f'{IC[code]}{r}', vals[code], FONT, F_INPUT if code in ('fuel', 'sub', 'grow') else None, fmt)
    ws.freeze_panes = 'D5'
    ws.column_dimensions['A'].width = 10
    for j in range(2, len(INPUT_COLS) + 1):
        ws.column_dimensions[L(j)].width = 10
    ws.row_dimensions[4].height = 36


# ---------------------------------------------------------------- Mitigation
def inputs_lookup(code, r):
    return f'INDEX(Inputs!${IC[code]}${IN_R0}:${IC[code]}${IN_R1},MATCH($B{r}&"|"&$C{r},Inputs!$A${IN_R0}:$A${IN_R1},0))'


def calc_formula(block, r, col, prev, use_lambda):
    """Calculation cell for a projection year (plain formula, or the named LAMBDA in the last year)."""
    if block == 'pre':
        return f'=PRETAX({prev}{r},$D{r})' if use_lambda else f'={prev}{r}*(1+$D{r})'
    if block == 'tax':
        return (f'=TAX($D{r},{col}${R_CP},$E{r},$F{r})' if use_lambda
                else f'=$D{r}+{col}${R_CP}*$E{r}*$F{r}')
    if block == 'post':
        a, b = f'{col}{r - 2 * BLOCK_H}', f'{col}{r - BLOCK_H}'
        return f'=POSTTAX({a},{b})' if use_lambda else f'={a}+{b}'
    if block == 'use':
        pn, pp = f'{col}{r - BLOCK_H}', f'{prev}{r - BLOCK_H}'
        if use_lambda:
            return f'=FUELUSE({prev}{r},{pn},{pp},{col}${R_GDP},$D{r},$E{r},$F{r},$G{r})'
        ratio = f'({pn}/{pp})'
        return (f'={prev}{r}*(1/(1+$G{r}))^(1+$E{r})*(1+{col}${R_GDP})^$D{r}'
                f'*{ratio}^$E{r}*{ratio}^($F{r}*(1+$E{r}))')


def base_formula(block, r, col):
    """Base-year cell: a lookup (data step) for levels; tax and post-tax use the plain formula."""
    if block == 'pre':
        return '=' + inputs_lookup('p0', r)
    if block == 'use':
        return '=' + inputs_lookup('f0', r)
    return calc_formula(block, r, col, None, False)


def group_cols(g):
    base = COL_G0 + (g - 1) * GW
    return base, list(range(base + 1, base + NY))


def build_mitigation(wb):
    ws = wb.create_sheet('Mitigation', 1)
    groups = [s[0] for s in SCENARIOS]
    last = COL_G0 + len(groups) * GW - 1
    for j, h in enumerate(['', 'Fuel', 'Sector'], 1):
        put(ws, f'{L(j)}1', h, FONT_B)
    for col in PARAM_COLS:
        put(ws, f'{col}1', f'Param {PARAM_COLS.index(col) + 1}', FONT_B)
    for col, h in DESC.items():
        put(ws, f'{col}1', h, FONT_B)
    title(ws, 'Mitigation module - price to fuel use (copy-pasteable, design 2)', last)
    band(ws, 3, 'Scenario assumptions (one column group per scenario; copy a whole group to add a scenario)', last)
    labels = {R_YEAR: ('Year', 'year', ''), R_NAME: ('Scenario name', '', 'Scenario input'),
              R_GDP: ('Real GDP growth', '%', 'WEO (lookup by country)'),
              R_CP: ('Carbon price', '$/tCO2', 'Scenario input')}
    for r, (lab, unit, src) in labels.items():
        put(ws, f'H{r}', lab, FONT_B)
        put(ws, f'I{r}', unit)
        put(ws, f'J{r}', src)
    put(ws, f'K{R_GDP}', 'mit.gdp.pos.pct', fill_=F_CODE)
    put(ws, f'K{R_CP}', 'mit.cp', fill_=F_CODE)

    for g, gname, cp in SCENARIOS:
        base, years = group_cols(g)
        put(ws, f'{L(base)}1', f'Scenario {g}', FONT_B)
        put(ws, f'{L(base)}{R_NAME}', gname, FONT_B, F_INPUT)
        for c in [base] + years:
            col, prev = L(c), L(c - 1)
            put(ws, f'{col}{R_YEAR}', '=Settings!$C$7' if c == base else f'={prev}{R_YEAR}+1', FONT_B,
                F_BASE if c == base else None, '0')
            put(ws, f'{col}{R_GDP}', (f'=INDEX(WEO!$F$4:$AZ$1000,MATCH(Settings!$C$4&"|gdp_growth",WEO!$A$4:$A$1000,0),'
                                      f'MATCH({col}${R_YEAR},WEO!$F$3:$AZ$3,0))'),
                fill_=F_BASE if c == base else None, fmt='0.0%')
            year = BASE_YEAR if c == base else YEARS[c - years[0]]
            put(ws, f'{col}{R_CP}', cp.get(year, 0), FONT, F_INPUT, '0.0')

    for block, (b0, text, code, unit, short, params) in BLOCKS.items():
        band(ws, b0, text, last)
        for k, pc in enumerate(PARAM_COLS):
            put(ws, f'{pc}{b0}', PARAM_LABELS[params[k]] if k < len(params) else '', FONT_B, F_BAND)
        src = {'pre': 'Inputs (base year), then growth', 'tax': 'Inputs', 'post': 'Calculation',
               'use': 'Inputs (base year), CPAT eq. 3.3.3'}[block]
        for i, (s, f) in enumerate(PAIRS):
            r = b0 + 1 + i
            put(ws, f'B{r}', f)
            put(ws, f'C{r}', s)
            for k, pc in enumerate(PARAM_COLS):
                if k < len(params):
                    put(ws, f'{pc}{r}', '=' + inputs_lookup(params[k], r), fmt='0.000')
                else:
                    put(ws, f'{pc}{r}', None, fill_=F_UNUSED)
            put(ws, f'H{r}', f'{short} | {SNAMES[s]} | {FNAMES[f]}')
            put(ws, f'I{r}', unit)
            put(ws, f'J{r}', src)
            put(ws, f'K{r}', f'mit.{code}.{f}.{s}', fill_=F_CODE)
            fmt = '#,##0.0' if block == 'use' else '0.000'
            for g in groups:
                base, years = group_cols(g)
                put(ws, f'{L(base)}{r}', base_formula(block, r, L(base)), fill_=F_BASE, fmt=fmt)
                for c in years:
                    lam = c == years[-1]
                    put(ws, f'{L(c)}{r}', calc_formula(block, r, L(c), L(c - 1), lam),
                        FONT_LAMBDA if lam else FONT, F_LAMBDA if lam else None, fmt)

    # Totals (fuel use)
    band(ws, B_TOT, 'Totals - fuel use (ktoe)', last)
    u0, u1 = B_USE + 1, B_USE + NROW
    r_sub0, r_fuel0 = B_TOT + 1, B_TOT + NS + 2
    r_tot = r_fuel0 + NF + 1
    rows = ([(r_sub0 + i, 'all', s[0], f'Total fuel use | {s[1]}', 'C') for i, s in enumerate(SUBSECTORS)]
            + [(r_fuel0 + i, f[0], 'all', f'Total fuel use | {f[1]}', 'B') for i, f in enumerate(FUELS)])
    for r, fu, sub, desc, keycol in rows:
        put(ws, f'B{r}', fu)
        put(ws, f'C{r}', sub)
        put(ws, f'H{r}', desc)
        put(ws, f'I{r}', 'ktoe')
        put(ws, f'K{r}', f'mit.enc.{fu}.{sub}', fill_=F_CODE)
        for g in groups:
            base, years = group_cols(g)
            for c in [base] + years:
                col = L(c)
                put(ws, f'{col}{r}', f'=SUMIF(${keycol}${u0}:${keycol}${u1},${keycol}{r},{col}${u0}:{col}${u1})',
                    fill_=F_BASE if c == base else None, fmt='#,##0.0')
    far = L(last + 10 * GW)
    extra = [(r_tot, 'Total fuel use (all subsectors, all fuels)', 'ktoe', 'mit.enc.all.all',
              lambda col: f'=SUM({col}{r_sub0}:{col}{r_sub0 + NS - 1})', '#,##0.0', None),
             (r_tot + 1, 'Check: sum by fuel - sum by subsector (should be 0)', 'ktoe', '',
              lambda col: f'=SUM({col}{r_fuel0}:{col}{r_fuel0 + NF - 1})-{col}{r_tot}', '0.000', F_CHECK),
             (r_tot + 3, 'Total fuel use, scenario 1 (same year)', 'ktoe', '',
              lambda col: (f'=INDEX(${L(COL_G0)}{r_tot}:${far}{r_tot},'
                           f'MATCH({col}${R_YEAR},${L(COL_G0)}${R_YEAR}:${far}${R_YEAR},0))'),
              '#,##0.0', None),
             (r_tot + 4, 'Change in total fuel use vs scenario 1', '%', 'mit.enc.all.all.pct',
              lambda col: f'={col}{r_tot}/{col}{r_tot + 3}-1', '0.00%', None)]
    for r, desc, unit, code, f, fmt, fl in extra:
        put(ws, f'H{r}', desc, FONT_B)
        put(ws, f'I{r}', unit)
        if code:
            put(ws, f'K{r}', code, fill_=F_CODE)
        for g in groups:
            base, years = group_cols(g)
            for c in [base] + years:
                put(ws, f'{L(c)}{r}', f(L(c)), fill_=fl or (F_BASE if c == base else None), fmt=fmt)

    # Layout: parameter columns D:G grouped and hidden by default.
    ws.column_dimensions.group('D', 'G', outline_level=1, hidden=True)
    ws.sheet_format.outlineLevelCol = 1
    ws.freeze_panes = f'{L(COL_G0)}9'
    for col, w in {'A': 3, 'B': 5, 'C': 5, 'H': 38, 'I': 7, 'J': 16, 'K': 15}.items():
        ws.column_dimensions[col].width = w
    for g in groups:
        base, years = group_cols(g)
        for c in [base] + years:
            ws.column_dimensions[L(c)].width = 9
        ws.column_dimensions[L(base + NY)].width = 3
    return last, r_tot


# ---------------------------------------------------------------- ReadMe
def build_readme(wb, r_tot):
    ws = wb.active
    ws.title = 'ReadMe'
    title(ws, f'CPAT mitigation module - copy-pasteable prototype v{VERSION}', 3)
    base2, years2 = group_cols(2)
    nxt = base2 + GW
    lam_lines = '; '.join(f'{n}({", ".join(p)}) = {b}' for n, (p, b) in LAMBDAS.items())
    lines = [
        ('Purpose', 'Auditable, copy-pasteable replacement of the CPAT mitigation price -> fuel-use chain. Each '
                    'block has one formula across all its cells, and a whole scenario group copies to a new '
                    'scenario and keeps working.'),
        ('Status', f'v{VERSION} prototype. Price -> fuel use only (no emissions, power or revenue yet).'),
        ('Layout (design 2)', f'Columns A:C section, fuel, subsector; D:G parameters (hidden by default - '
                              f'click + above column H to show); H:K description, unit, source, output code; then '
                              f'one group per scenario: base year {BASE_YEAR} in column L (as in legacy '
                              f'Mitigation), {YEARS[0]}-{YEARS[-1]}, one gap column.'),
        ('Nesting', 'Across: scenario (outer) > year (inner). Down: variable block (outer) > subsector > fuel '
                    f'(inner). {NS} CPAT energy-use subsectors x {NF} fuels = {NROW} rows per block.'),
        ('Blocks', f'Pre-tax price (row {B_PRE}), Tax (row {B_TAX}), Post-tax price (row {B_POST}), '
                   f'Fuel use (row {B_USE}), Totals (row {B_TOT}).'),
        ('Data step vs formulas', 'All lookups happen in sheet Inputs (one row per fuel|subsector) and in '
                                  'Mitigation D:G and the base-year column, each a single INDEX/MATCH on the '
                                  'fuel|subsector key. Calculation cells contain no lookups.'),
        ('Equations', 'Pre-tax: p(t) = p(t-1) x (1 + growth). Tax: base tax + carbon price x EF x coverage '
                      '(base tax = retail - supply cost, so it includes existing excise, subsidies and VAT). '
                      'Post-tax = pre-tax + tax. Fuel use: F(t) = F(t-1) x (1/(1+alpha))^(1+eps_U) x '
                      '(1+g)^eps_Y x (p(t)/p(t-1))^eps_U x (p(t)/p(t-1))^(eps_F(1+eps_U)) (CPAT documentation '
                      '3.3.3, without the Covid factor and shadow price).'),
        ('LAMBDA column', f'{YEARS[-1]} (orange, dark-red text) calls named LAMBDAs with the same equations; '
                          f'{YEARS[0]}-{YEARS[-2]} are plain formulas. Drag the {YEARS[-1]} cell back over the '
                          f'row to use the LAMBDA everywhere, or drag {YEARS[-2]} forward to remove it. '
                          f'Names (Formulas > Name Manager): {lam_lines}.'),
        ('Parameters', 'Global, not scenario-specific: scenarios differ only in their assumption rows '
                       f'(row {R_CP} carbon price). Pre-tax price growth is an input in Inputs (0 % now).'),
        ('Lookups by country', 'Settings!C4 selects the country. Elasticities (via income group) and prices '
                               'cover all CPAT countries; base-year energy use, GDP growth and EFs are Egypt only '
                               'so far (other countries return #N/A until their rows are added).'),
        ('Add a scenario', f'Copy columns {L(base2)}:{L(base2 + GW - 1)} (a whole group with its gap column) and '
                           f'paste at column {L(nxt)}; rename it in row {R_NAME} and edit the carbon prices in '
                           f'row {R_CP}.'),
        ('Colours', 'Green = input; tan = base-year column; light blue = codes; grey = unused parameter slot; '
                    'orange = LAMBDA column; yellow = check (NORMS section 2).'),
        ('Deviations from NORMS 1.1', 'Parameter columns D:G sit between Sector and Description, so Description '
                                      '.. Output Code are in H:K and Input Code, Note and Helper are dropped; the '
                                      'base year 2022 stays in column L as in legacy.'),
        ('Data', 'Data sheets: Countries, Elasticities, Prices_dom, EnergyCons, WEO, EF_GHG; mappings in '
                 'Mapping. Built by build_v0_2.py from data/*.csv (extract_data_v0_1.py).'),
        ('Checks', f'Row {r_tot + 1} (sum by fuel - sum by subsector) should be 0. check_v0_2.py tests formula '
                   'uniformity, the LAMBDA encoding and equivalence, the scenario-copy behaviour and an '
                   'independent recomputation.'),
    ]
    for i, (k, v) in enumerate(lines, 4):
        put(ws, f'A{i}', k, FONT_B)
        c = put(ws, f'B{i}', v)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws[f'A{i}'].alignment = Alignment(vertical='top')
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 120


def main():
    wb = Workbook()
    build_settings(wb)
    build_mapping(wb)
    build_inputs(wb)
    last, r_tot = build_mitigation(wb)
    div = wb.create_sheet('DATA->')
    put(div, 'A1', 'Data sheets follow', FONT_B)
    build_data_sheets(wb)
    build_readme(wb, r_tot)
    for name in LAMBDAS:
        wb.defined_names[name] = DefinedName(name, attr_text=lambda_xml(name), comment=LAMBDA_NOTES[name])
    order = ['ReadMe', 'Settings', 'Mitigation', 'Inputs', 'Mapping', 'DATA->', 'Countries', 'Elasticities',
             'Prices_dom', 'EnergyCons', 'WEO', 'EF_GHG']
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(OUT)
    print('Saved', OUT)


if __name__ == '__main__':
    main()
