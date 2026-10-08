"""Checks for CPAT_Mitigation_CopyPaste_v0.1.xlsx. Needs LibreOffice (soffice) to recalculate.

1. Recalculate; no error values on Mitigation.
2. Formula uniformity: within each block, every projection cell (all rows, all scenario groups) has the
   same formula in relative R1C1 form; likewise every base-year cell.
3. Copy test: copy scenario group 2 to group 3 (as a user would: paste the whole group, type scenario
   number 3). Group 3 must equal group 2. Copy group 2 to group 4 with the carbon price set to 0: group 4
   must equal group 1.
4. Independent recomputation in Python from data/*.csv; compare every block cell.
5. Sanity: scenarios 1 and 2 identical before 2027; uncovered (oen) and zero-EF (bio) rows identical.

    python check_v0_1.py   ->  writes check_report_v0_1.md
"""
import csv
import os
import re
import shutil
import subprocess
import tempfile

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

import build_v0_1 as B

HERE = os.path.dirname(os.path.abspath(__file__))
WB = B.OUT
REF = re.compile(r"(?<![A-Za-z_\d.\"])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(])")
TOL = 1e-9
report = []


def log(line=''):
    print(line)
    report.append(line)


def recalc(path):
    out = tempfile.mkdtemp()
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out, path],
                   check=True, capture_output=True, timeout=600)
    return os.path.join(out, os.path.basename(path))


def r1c1(formula, row, col):
    def sub(m):
        ca, cl, ra, rw = m.group(1), m.group(2), m.group(3), int(m.group(4))
        c = CI(cl)
        cs = f'C{c}' if ca else f'C[{c - col}]'
        rs = f'R{rw}' if ra else f'R[{rw - row}]'
        return rs + cs
    return REF.sub(sub, formula)


def block_rows(b0):
    return range(b0 + 1, b0 + 1 + B.NROW)


def groups_of(ws_or_n):
    return [B.group_cols(g) for g in range(1, ws_or_n + 1)]


# ------------------------------------------------------------------ 2. uniformity
def check_uniformity(wbf):
    ws = wbf['Mitigation']
    ok = True
    sections = {name: block_rows(b0) for name, (b0, *_rest) in B.BLOCKS.items()}
    r_sub0 = B.B_TOT + 1
    sections['tot_sub'] = range(r_sub0, r_sub0 + B.NS)
    sections['tot_fuel'] = range(B.B_TOT + B.NS + 2, B.B_TOT + B.NS + 2 + B.NF)
    sections['macro_gdp'] = [B.R_GDP]
    for name, rows in sections.items():
        base_set, proj_set = set(), set()
        for _, base, years in groups_of(len(B.SCENARIOS)):
            for r in rows:
                base_set.add(r1c1(ws.cell(r, base).value, r, base))
                for c in years:
                    proj_set.add(r1c1(ws.cell(r, c).value, r, c))
        same = len(proj_set) == 1 and len(base_set) == 1
        ok &= same
        log(f'- `{name}`: {len(rows)} rows x {len(B.SCENARIOS)} groups -> {len(base_set)} base-year formula, '
            f'{len(proj_set)} projection formula {"OK" if same else "FAIL"}')
        if name in ('use', 'tax'):
            log(f'  - projection R1C1: `{next(iter(proj_set))[:200]}...`')
    # Parameter columns: one formula per (block, slot) column across all rows.
    for name, (b0, *_rest) in B.BLOCKS.items():
        for j in range(B.COL_P0, B.COL_P1 + 1):
            forms = {r1c1(str(ws.cell(r, j).value), r, j) for r in block_rows(b0)}
            if len(forms) != 1:
                ok = False
                log(f'- parameter column {L(j)} in block {name}: {len(forms)} formulas FAIL')
    log(f'- parameter columns {L(B.COL_P0)}:{L(B.COL_P1)}: one formula (or input/blank) per block and column '
        f'{"OK" if ok else "see above"}')
    return ok


# ------------------------------------------------------------------ 3. copy test
def copy_group(ws, src_g, dst_g, last_row, cp_zero=False):
    s0, _, _ = B.group_cols(src_g)
    d0, _, _ = B.group_cols(dst_g)
    for k in range(B.GW):
        sc, dc = s0 + k, d0 + k
        for r in range(1, last_row + 1):
            v = ws.cell(r, sc).value
            if isinstance(v, str) and v.startswith('='):
                v = Translator(v, origin=f'{L(sc)}{r}').translate_formula(f'{L(dc)}{r}')
            ws.cell(r, dc).value = v
    ws.cell(B.R_SCEN, d0).value = dst_g          # the user types the new scenario number
    if cp_zero:
        for c in range(d0 + 1, d0 + B.GW):
            ws.cell(B.R_CP, c).value = 0


def check_copy(path):
    wbf = openpyxl.load_workbook(path)
    ws = wbf['Mitigation']
    last_row = ws.max_row
    copy_group(ws, 2, 3, last_row)
    copy_group(ws, 2, 4, last_row, cp_zero=True)
    tmp = os.path.join(tempfile.mkdtemp(), 'copytest.xlsx')
    wbf.save(tmp)
    ws = openpyxl.load_workbook(recalc(tmp), data_only=True)['Mitigation']
    ok = True
    for a, b, label in [(2, 3, 'group 3 (copy of 2) = group 2'), (1, 4, 'group 4 (copy of 2, carbon price 0) = group 1')]:
        _, ba, ya = B.group_cols(a)
        _, bb, yb = B.group_cols(b)
        diff, n = 0.0, 0
        for r in range(B.R_YEAR, last_row + 1):
            if r == B.R_CP:
                continue
            for ca, cb in zip([ba] + ya, [bb] + yb):
                va, vb = ws.cell(r, ca).value, ws.cell(r, cb).value
                if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                    if r in (B.R_SCEN,):
                        continue
                    diff = max(diff, abs(va - vb))
                    n += 1
                elif va != vb and r != B.R_NAME:
                    diff = float('inf')
        res = diff < 1e-9
        ok &= res
        log(f'- {label}: {n} numeric cells, max abs diff {diff:.3g} {"OK" if res else "FAIL"}')
    return ok


# ------------------------------------------------------------------ 4. independent recomputation
def read(name):
    with open(os.path.join(HERE, 'data', name), encoding='utf-8') as f:
        return list(csv.DictReader(f))


def recompute():
    units = {r['fuel']: float(r['gj_per_price_unit']) for r in read('fuel_units.csv')}
    prices = {r['country_year']: r for r in read('prices_dom.csv')}['EGY2022']
    ef = {r['key']: float(r['ef_tco2_per_gj']) for r in read('ef_co2.csv')}
    el = {r['key']: float(r['LMIC']) for r in read('elasticities.csv')}
    use0 = {(r['subsector'], r['fuel']): float(r['ktoe']) for r in read('energy_use.csv')}
    weo = read('weo.csv')[0]
    sub = {s[0]: s for s in B.SUBSECTORS}
    fue = {f[0]: f for f in B.FUELS}
    res = {}
    for g, _, cp in B.SCENARIOS:
        for s, *_ in B.SUBSECTORS:
            for f, *_ in B.FUELS:
                _, _, grp, es, efs, cgs, cov = sub[s]
                _, _, elf, aef, rule, _ = fue[f]
                code = f'{f}.{cgs if rule == "sub" else "all"}'
                sp = float(prices[f'mit.sp.{code}']) / units[f]
                btax = (float(prices[f'mit.rp.{code}']) - float(prices[f'mit.sp.{code}'])) / units[f]
                e = ef[f'EGY|{f}|{efs}']
                eY, eU, eF = el[f'inc|{elf}|{es}'], el[f'usg|{elf}|{es}'], el[f'eff|{elf}|{es}']
                a = el[f'aei|{aef}|{es}']
                pre, tax, post, use = [sp], [btax], [sp + btax], [use0[(s, f)]]
                for y in B.YEARS:
                    pre.append(pre[-1])
                    tax.append(btax + cp.get(y, 0) * e * cov)
                    post.append(pre[-1] + tax[-1])
                    ratio = post[-1] / post[-2]
                    g_ = float(weo[str(y)])
                    use.append(use[-1] * (1 / (1 + a)) ** (1 + eU) * (1 + g_) ** eY * ratio ** eU
                               * ratio ** (eF * (1 + eU)))
                for blk, vals in (('pre', pre), ('tax', tax), ('post', post), ('use', use)):
                    res[(g, blk, s, f)] = vals
    return res


def check_values(ws):
    exp = recompute()
    ok, worst = True, (0.0, None)
    for name, (b0, *_rest) in B.BLOCKS.items():
        for i, r in enumerate(block_rows(b0)):
            s, f = B.SUBSECTORS[i // B.NF][0], B.FUELS[i % B.NF][0]
            for g, *_ in B.SCENARIOS:
                _, base, years = B.group_cols(g)
                for k, c in enumerate([base] + years):
                    v, e = ws.cell(r, c).value, exp[(g, name, s, f)][k]
                    d = abs(v - e) / max(1.0, abs(e))
                    if d > worst[0]:
                        worst = (d, f'{L(c)}{r}')
    ok = worst[0] < TOL
    log(f'- all {len(B.BLOCKS)} blocks x {B.NROW} rows x {len(B.SCENARIOS)} groups x {1 + len(B.YEARS)} years: '
        f'max relative diff {worst[0]:.2e} (at {worst[1]}) {"OK" if ok else "FAIL"}')
    return ok


# ------------------------------------------------------------------ 5. sanity and results
def check_sanity(ws):
    ok = True
    _, b1, y1 = B.group_cols(1)
    _, b2, y2 = B.group_cols(2)
    cols1, cols2 = [b1] + y1, [b2] + y2
    years = [B.BASE_YEAR] + B.YEARS
    pre = max(abs(ws.cell(r, c1).value - ws.cell(r, c2).value)
              for r in block_rows(B.B_USE) for c1, c2, y in zip(cols1, cols2, years) if y < 2027)
    log(f'- scenario 1 vs 2 before 2027 (fuel use): max abs diff {pre:.3g} {"OK" if pre < TOL else "FAIL"}')
    ok &= pre < TOL
    for label, pick in [('oen rows (no coverage)', lambda s, f: s == 'oen'), ('bio rows (EF 0)', lambda s, f: f == 'bio')]:
        d = max(abs(ws.cell(r, c1).value - ws.cell(r, c2).value)
                for i, r in enumerate(block_rows(B.B_USE)) if pick(B.SUBSECTORS[i // B.NF][0], B.FUELS[i % B.NF][0])
                for c1, c2 in zip(cols1, cols2))
        log(f'- {label}: scenario 1 = scenario 2, max abs diff {d:.3g} {"OK" if d < TOL else "FAIL"}')
        ok &= d < TOL
    r_tot = B.B_TOT + B.NS + 2 + B.NF + 1
    chk = max(abs(ws.cell(r_tot + 1, c).value) for c in cols1 + cols2)
    log(f'- check row {r_tot + 1} (sum by fuel - sum by subsector): max abs {chk:.3g} {"OK" if chk < 1e-6 else "FAIL"}')
    ok &= chk < 1e-6
    log('')
    log('Results (total fuel use, ktoe, and change vs scenario 1):')
    log('')
    log('| Year | Scenario 1 | Scenario 2 | Change |')
    log('|---|---|---|---|')
    for c1, c2, y in zip(cols1, cols2, years):
        if y in (2022, 2026, 2027, 2030, 2035):
            a, b = ws.cell(r_tot, c1).value, ws.cell(r_tot, c2).value
            log(f'| {y} | {a:,.0f} | {b:,.0f} | {b / a - 1:+.1%} |')
    return ok


def main():
    log(f'# Check report - CPAT_Mitigation_CopyPaste_v{B.VERSION}')
    log('')
    calc = recalc(WB)
    wbv = openpyxl.load_workbook(calc, data_only=True)
    ws = wbv['Mitigation']
    errs = [c.coordinate for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and (c.value.startswith('#') or c.value.startswith('Err:'))]
    log('## 1. Recalculation (LibreOffice)')
    log(f'- error values on Mitigation: {len(errs)} {"OK" if not errs else "FAIL " + str(errs[:10])}')
    log('')
    log('## 2. Formula uniformity (relative R1C1)')
    u = check_uniformity(openpyxl.load_workbook(WB))
    log('')
    log('## 3. Scenario-group copy test')
    cp = check_copy(WB)
    log('')
    log('## 4. Independent recomputation (Python, from data/*.csv)')
    v = check_values(ws)
    log('')
    log('## 5. Sanity')
    s = check_sanity(ws)
    log('')
    allok = not errs and u and cp and v and s
    log(f'**Overall: {"PASS" if allok else "FAIL"}**')
    with open(os.path.join(HERE, f'check_report_v{B.VERSION}.md'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(report) + '\n')
    shutil.rmtree(os.path.dirname(calc), ignore_errors=True)


if __name__ == '__main__':
    main()
