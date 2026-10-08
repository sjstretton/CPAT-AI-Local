"""Extract the source data for CPAT_Mitigation_CopyPaste_v0.1 into data/*.csv.

Sources (all non-proprietary, already in the repo):
- Legacy CPAT `cpat_excel_original/CPAT 1.0pre_456_NoPropData.xlsb` (read-only):
    * `Elasticities` rows 138-141 (country list, region, income group),
      rows 143-174 (income), 178-209 (usage), 211-242 (efficiency),
      rows 256-282 (autonomous efficiency improvement), income-group columns LIC..HIC.
    * `Mitigation` rows 980-1009 (IIASA CO2 EFs, Egypt, tCO2/GJ, and volume-unit-to-GJ factors).
- `cpat_excel_new/standalone_initial_prototypes/CPAT_PriceModule_Data_1.xlsx`, sheet `DomesticPrices`
  (221 countries, 2021-2024).
- `cpat_excel_new/standalone_working_version/CPAT_Industry_Kernel_Egypt_v1.6.xlsx`:
    * `Data_Energy` (base-year energy use, Egypt, ktoe; from legacy Mitigation rows 583-602),
    * `Data_Macro` (real GDP growth, Egypt, scenario 1; legacy Mitigation row 2446).

Needs pyxlsb (pip install pyxlsb). Run from the repo root or anywhere:
    python extract_data_v0_1.py
"""
import csv
import os

import openpyxl
from pyxlsb import open_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
LEGACY = os.path.join(ROOT, 'cpat_excel_original', 'CPAT 1.0pre_456_NoPropData.xlsb')
PRICES = os.path.join(ROOT, 'cpat_excel_new', 'standalone_initial_prototypes', 'CPAT_PriceModule_Data_1.xlsx')
KERNEL = os.path.join(ROOT, 'cpat_excel_new', 'standalone_working_version', 'CPAT_Industry_Kernel_Egypt_v1.6.xlsx')
OUT = os.path.join(HERE, 'data')

GROUPS = ['LIC', 'LMIC', 'UMIC', 'HIC']
SECTOR_CODE = {'Transport': 'tra', 'Residential': 'res', 'Industry': 'ind', 'Services': 'srv'}
FUEL_CODE = {'Coal': 'coa', 'Natural Gas': 'nga', 'Gasoline': 'gso', 'Diesel': 'die',
             'Oil Products': 'oil', 'Oil': 'oil', 'Biomass': 'bio', 'Renewables': 'ren',
             'Electricity': 'ecy'}
SUBSECTORS = ['rod', 'ral', 'avi', 'nav', 'res', 'foo', 'srv',
              'mch', 'irn', 'nfm', 'mac', 'cem', 'omn', 'cst', 'ftr', 'oen']
FUELS = ['coa', 'nga', 'gso', 'die', 'lpg', 'ker', 'oop', 'bio']


def write(name, header, rows):
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, name), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f'{name}: {len(rows)} rows')


def legacy_sheet(name, ncols=None):
    with open_workbook(LEGACY) as wb:
        with wb.get_sheet(name) as sh:
            return [[c.v for c in r][:ncols] if ncols else [c.v for c in r] for r in sh.rows()]


def extract_elasticities():
    rows = legacy_sheet('Elasticities')
    # Country header block, rows 138-141 (1-based); countries start at column index 13.
    names, codes, regions, groups = (rows[i] for i in (137, 138, 139, 140))
    countries = []
    for j in range(13, len(codes)):
        if codes[j] not in (None, ''):
            countries.append([codes[j], names[j], regions[j], groups[j], f'Elasticities!{j}:138-141'])
    write('countries.csv', ['code', 'name', 'region', 'income_group', 'source'], countries)

    out = []
    # (type code, first row, last row, column index of LIC); LIC..HIC are 4 consecutive columns.
    blocks = [('inc', 143, 174, 9), ('usg', 178, 209, 9), ('eff', 211, 242, 9)]
    for typ, r0, r1, c0 in blocks:
        for i in range(r0 - 1, r1):
            r = rows[i]
            fuel, sector = FUEL_CODE[r[3]], SECTOR_CODE[r[4]]
            out.append([f'{typ}|{fuel}|{sector}', typ, fuel, sector, r[5]]
                       + [r[c0 + k] for k in range(4)] + [f'Elasticities row {i + 1}'])
    # Autonomous efficiency improvement: LIC..HIC start one column earlier (index 8).
    for i in range(255, 282):
        r = rows[i]
        if r[2] != 'Autonomous efficiency':
            continue
        fuel, sector = FUEL_CODE[r[3]], SECTOR_CODE[r[4]]
        out.append([f'aei|{fuel}|{sector}', 'aei', fuel, sector, r[5]]
                   + [r[8 + k] for k in range(4)] + [f'Elasticities row {i + 1}'])
    write('elasticities.csv', ['key', 'type', 'fuel_group', 'sector_group', 'legacy_code'] + GROUPS + ['source'], out)


def extract_ef():
    rows = legacy_sheet('Mitigation', 20)
    out, conv = [], {}
    for i in range(979, 1009):  # rows 980-1009
        r = rows[i]
        fuel, sec = r[1], r[2]
        if fuel not in FUELS:
            continue
        conv.setdefault(fuel, (r[8], r[9]))
        secs = ['pow', 'rod', 'res', 'ind'] if sec == 'all' else [sec]
        for s in secs:  # sector-specific rows come first and win over 'all' in the de-duplication below
            out.append([f'EGY|{fuel}|{s}', 'EGY', fuel, s, r[15], f'Mitigation row {i + 1} ({r[7]})'])
    seen, uniq = set(), []
    for row in out:
        if row[0] not in seen:
            seen.add(row[0])
            uniq.append(row)
    write('ef_co2.csv', ['key', 'country', 'fuel', 'ef_sector', 'ef_tco2_per_gj', 'source'], uniq)
    # Price-unit to GJ factor per fuel (legacy 'Conversion factor (volume unit to GJ)').
    write('fuel_units.csv', ['fuel', 'gj_per_price_unit', 'ef_unit_in_legacy'],
          [[f, conv[f][0], conv[f][1]] for f in FUELS])


def extract_prices():
    ws = openpyxl.load_workbook(PRICES, read_only=True, data_only=True)['DomesticPrices']
    rows = list(ws.iter_rows(values_only=True))
    # Keep the legacy row-banded header (rows 1-4), code row 5, data from row 6.
    write('prices_dom.csv', list(rows[4]), [list(r) for r in rows[5:] if r[0]])
    write('prices_dom_header.csv', ['band'] + [f'c{j}' for j in range(len(rows[0]) - 1)],
          [list(r) for r in rows[:4]])


def extract_kernel():
    wb = openpyxl.load_workbook(KERNEL, read_only=True, data_only=True)
    rows = list(wb['Data_Energy'].iter_rows(values_only=True))
    hdr = rows[1]
    col = {h: j for j, h in enumerate(hdr) if h}
    out = []
    for r in rows[2:]:
        fuel = r[2]
        if fuel not in FUELS:
            continue
        for s in SUBSECTORS:
            out.append([f'EGY|2022|{s}|{fuel}', 'EGY', 2022, s, fuel, r[col[s]] or 0.0,
                        'Kernel v1.6 Data_Energy = legacy Mitigation rows 583-602'])
    write('energy_use.csv', ['key', 'country', 'year', 'subsector', 'fuel', 'ktoe', 'source'], out)

    rows = list(wb['Data_Macro'].iter_rows(values_only=True))
    ycols = [(j, int(y)) for j, y in enumerate(rows[0]) if isinstance(y, (int, float))]
    g = next(r for r in rows if r[0] == 'egy.mit.gdp.pos.pct.1')
    write('weo.csv', ['key', 'country', 'variable', 'unit', 'source'] + [y for _, y in ycols],
          [['EGY|gdp_growth', 'EGY', 'Real GDP growth', 'fraction',
            'Kernel v1.6 Data_Macro egy.mit.gdp.pos.pct.1 = legacy Mitigation row 2446']
           + [g[j] for j, _ in ycols]])


if __name__ == '__main__':
    extract_elasticities()
    extract_ef()
    extract_prices()
    extract_kernel()
