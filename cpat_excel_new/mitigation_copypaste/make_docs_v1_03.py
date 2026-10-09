"""Word documentation of CPAT-AI-Mitigation-MVP v1.03, assembled from the Markdown notes (pandoc), with appendices
read from the shipped workbook (differences with legacy CPAT from sheet LegacyDiff, version log from Settings) and
from the check report (verification summary).

    python make_docs_v1_03.py  ->  CPAT-AI-Mitigation-MVP_Documentation_v1.03.docx
"""
import os
import re
import subprocess
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'CPAT-AI-Mitigation-MVP_Documentation_v1.03.docx')
WORKBOOK = os.path.join(HERE, 'CPAT-AI-Mitigation-MVP-v1.03.xlsm')
REPORT = os.path.join(HERE, 'check_report_v1.03.md')
PARTS = [('Model overview, layout and files', 'README.md'),
         ('Domestic energy prices', 'PriceProjection_Method_v0.4.md'),
         ('New ETS', 'ETS_Method_v0_1.md'),
         ('Multiple scenarios and stored results', 'Scenarios_Method_v0_1.md')]


def part(title, name):
    text = open(os.path.join(HERE, name), encoding='utf-8').read()
    text = re.sub(r'\A# .*\n', '', text)                      # drop the note's own title
    text = re.sub(r'^(#+) ', lambda m: '#' + m.group(1) + ' ', text, flags=re.M)   # demote one level
    return f'# {title}\n\n*Source: `{name}`.*\n\n{text}\n'


def cell(v):
    return str(v if v is not None else '').replace('|', '/').replace('\n', ' ')


def table(header, rows, widths=None):
    """Pipe table; widths = relative column widths (pandoc reads them from the separator dashes)."""
    widths = widths or [1] * len(header)
    out = ['| ' + ' | '.join(header) + ' |', '|' + '|'.join('-' * (6 * w) for w in widths) + '|']
    out += ['| ' + ' | '.join(cell(v) for v in r) + ' |' for r in rows]
    return '\n'.join(out) + '\n'


def appendices():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    ld = [r for r in wb['LegacyDiff'].iter_rows(min_row=5, max_col=7, values_only=True) if r[1]]
    log = [r[1:5] for r in wb['Settings'].iter_rows(min_row=23, max_col=5, values_only=True) if r[1]]
    text = ['# Appendix A. Differences with legacy CPAT', '',
            '*Source: workbook sheet `LegacyDiff`. Type: Decision = user decision; Assumption = data assumption; '
            'Simplification; Not yet; Data; Method; Same.*', '',
            table(['#', 'Area', 'Legacy CPAT', 'This model', 'Why', 'Effect', 'Type'], ld, [1, 4, 5, 6, 5, 4, 5]),
            '# Appendix B. Version log', '', '*Source: workbook sheet `Settings`.*', '',
            table(['Version', 'Date', 'Description', 'Max abs regression diff'], log, [2, 3, 12, 5]),
            '# Appendix C. Verification', '',
            '*Source: `check_report_v1.03.md` (`python check_v1_03.py`, LibreOffice recalculation). Every check '
            'section and its result:*', '']
    rep_lines = open(REPORT, encoding='utf-8').read().splitlines()
    sec, oks, fails, rows = None, 0, 0, []
    for ln in rep_lines + ['## end']:
        if ln.startswith('## '):
            if sec:
                rows.append([sec, oks, fails])
            sec, oks, fails = ln[3:], 0, 0
        elif ln.strip().startswith('- ') or ln.strip().startswith('  - '):
            oks += ln.rstrip().endswith('OK')
            fails += 'FAIL' in ln
    overall = next((ln for ln in rep_lines if ln.startswith('**Overall')), '')
    text += [table(['Check section', 'Checks OK', 'Checks failed'], rows[:-1] if rows[-1][0] == 'end' else rows, [8, 2, 2]),
             f'{overall}', '',
             'The checks recompute every value independently in Python from the source data; test the formula '
             'variants (LAMBDA expanded, plain formula dragged into the LAMBDA column, LAMBDA copied over whole rows); '
             'copy scenario groups as a user would; compare with the previous version on every shared output code; '
             'and run the embedded VBA macro in LibreOffice against the Python emulation.', '',
             '# Appendix D. Presentations', '',
             '- `presentation/CPAT-AI-Mitigation-MVP_Overview_v1.0.pptx`: 7-slide pitch and user guide (v1.00).',
             '- `presentation/CPAT-AI-Mitigation-MVP_AdvancedFeatures_v1.0.pptx`: 4 slides on v1.01-v1.03 (cap-based '
             'ETS, multiple scenarios with the batch macro, stored results and comparison).', '']
    return '\n'.join(text)


def main():
    head = ('---\ntitle: "CPAT-AI-Mitigation-MVP v1.03: documentation"\n'
            'subtitle: "AI-generated, copy-pasteable replacement of the CPAT mitigation module (Egypt data)"\n'
            'date: "2026-10-09"\n---\n\n'
            '# About this document\n\nWorkbook: `CPAT-AI-Mitigation-MVP-v1.03.xlsm` (macro-enabled) (built by `build_v1_03.py`, '
            'checked by `check_v1_03.py`, report `check_report_v1.03.md`: PASS). This document collects the '
            'model overview, the method notes, and appendices on the differences with legacy CPAT (sheet `LegacyDiff`), '
            'the version log (sheet `Settings`), the verification and the presentations. Caveats per task are in '
            '`CAVEATS.md` at the repository root.\n\n')
    md = head + '\n'.join(part(t, n) for t, n in PARTS) + '\n' + appendices()
    src = os.path.join(tempfile.mkdtemp(), 'doc.md')
    open(src, 'w', encoding='utf-8').write(md)
    subprocess.run(['pandoc', src, '-o', OUT, '--toc', '--toc-depth=2', '-f', 'markdown-yaml_metadata_block+yaml_metadata_block'],
                   check=True)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
