"""Checks for CPAT_Mitigation_CopyPaste_v0.7.xlsx. Needs LibreOffice (soffice) to recalculate.

LibreOffice 24.2 cannot evaluate LAMBDA, so the 2035 LAMBDA cells are checked by expansion (body with the
arguments substituted) and by dragging the plain 2034 formula forward.

1. Shipped file: errors only in the 2035 columns.
2. Uniformity (relative R1C1): one formula per variable for base year / plain years / 2035, across all
   subsectors, fuels and scenarios; one formula per D:G column per variable; Section-1 input rows; paths;
   label and code columns.
3. LAMBDA encoding.
4. Values vs an independent Python recomputation (both 2035 variants).
5. Scenario tests (MTInputs column + Mitigation group copied as a user would): copy = source; zero carbon price =
   baseline; legacy Egypt carbon-tax settings = legacy row 8582; fuel price reform and feebates vs Python;
   the feebate shadow price does not change fuel use (not yet used).
6. Regression vs v0.6 on every output code the two versions share.
7. Labels and codes; 8. outline; 9. MTInputs = template; 10. format; 11. sanity and results.

    python check_v0_7.py   ->  writes check_report_v0.7.md
"""
import csv
import importlib.util
import os
import re
import subprocess
import tempfile

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

import build_v0_7 as B

HERE = os.path.dirname(os.path.abspath(__file__))
REF = re.compile(r"(?<![A-Za-z_\d.\"])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(])")
CALL = re.compile(r'^=(' + '|'.join(B.LAMBDAS) + r')\((.*)\)$')
TOL = 1e-9
report = []


def log(line=''):
    print(line)
    report.append(line)


def recalc(path):
    out = tempfile.mkdtemp()
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out, path],
                   check=True, capture_output=True, timeout=900)
    return openpyxl.load_workbook(os.path.join(out, os.path.basename(path)), data_only=True)['Mitigation']


def r1c1(formula, row, col):
    def sub(m):
        c, rw = CI(m.group(2)), int(m.group(4))
        return (f'R{rw}' if m.group(3) else f'R[{rw - row}]') + (f'C{c}' if m.group(1) else f'C[{c - col}]')
    return REF.sub(sub, str(formula))


def is_err(v):
    return isinstance(v, str) and (v.startswith('#') or v.startswith('Err:'))


def expand(formula, lambdas=None):
    lambdas = lambdas or B.LAMBDAS
    m = re.match(r'^=(' + '|'.join(lambdas) + r')\((.*)\)$', formula)
    if not m:
        return formula
    params, body = lambdas[m.group(1)]
    args = m.group(2).split(',')
    assert len(args) == len(params), formula
    for prm in sorted(params, key=len, reverse=True):
        body = re.sub(rf'(?<![A-Za-z_.]){prm}(?![A-Za-z_\d(])', f'({args[params.index(prm)]})', body)
    return '=' + body


def groups_in(ws):
    return [g for g in range(1, 40) if ws.cell(B.R_SCEN, B.group_cols(g)[0]).value not in (None, '')]


def variant(wb, how):
    """how = 'expand' (LAMBDA bodies substituted in 2035) or 'drag' (plain 2034 formula dragged into 2035)."""
    ws = wb['Mitigation']
    for g in groups_in(ws):
        c = B.group_cols(g)[2][-1]
        for var in B.LAMBDA_VARS:
            for r, _s, _f in B.var_rows(var):
                if how == 'expand':
                    ws.cell(r, c).value = expand(ws.cell(r, c).value)
                else:
                    ws.cell(r, c).value = Translator(ws.cell(r, c - 1).value, origin=f'{L(c - 1)}{r}') \
                        .translate_formula(f'{L(c)}{r}')
    path = os.path.join(tempfile.mkdtemp(), f'{how}.xlsx')
    wb.save(path)
    return path


# ------------------------------------------------------------------ independent recomputation
def read(name):
    with open(os.path.join(HERE, 'data', name), encoding='utf-8') as f:
        return list(csv.DictReader(f))


def mt_used():
    tpl = openpyxl.load_workbook(B.MT_TEMPLATE)['MTInputs']
    return {r: tpl.cell(r, 6).value for r in range(8, 416)}


def recompute(mt):
    """mt: MTInputs row -> value for one scenario. Returns {(var, s, f): [values base..last]}."""
    units = {r['fuel']: float(r['gj_per_price_unit']) for r in read('fuel_units.csv')}
    prices = {r['country_year']: r for r in read('prices_dom.csv')}['EGY2022']
    ef = {r['key']: float(r['ef_tco2_per_gj']) for r in read('ef_co2.csv')}
    el = {r['key']: float(r['LMIC']) for r in read('elasticities.csv')}
    use0 = {(r['subsector'], r['fuel']): float(r['ktoe']) for r in read('energy_use.csv')}
    weo = read('weo.csv')[0]
    sub = {s[0]: s for s in B.SUBSECTORS}
    fue = {f[0]: f for f in B.FUELS}
    years = [B.BASE_YEAR] + B.YEARS
    n = B.MT_NAMES

    def path(y, y0, y1, s0, t, linear=False):
        if y < y0:
            return 0.0
        k = y - y0 if (y <= y1 or linear) else y1 - y0
        return s0 + (t - s0) / max(y1 - y0, 1) * k
    cp = {y: path(y, mt[n['CPIntro']], mt[n['CPOutro']], mt[n['CPLevelStart']], mt[n['CPLevelTarget']],
                  mt[n['ExtendCarbonPriceBeyondOutro']] == 'Linear*') for y in years}
    fpr = {pc: {y: path(y, mt[140], mt[141], mt[r0] or 0, mt[r1] or 0) for y in years}
           for pc, _l, r0, r1 in B.PRICE_FUELS}
    fb = {s: {y: path(y, mt[n['D_FeebateIntro']], mt[n['D_FeebateOutro']], mt[n[a]], mt[n[b]]) for y in years}
          for s, a, b in B.FEEBATE_SECTORS}
    res = {}
    for s, f in B.PAIRS:
        _, _, grp, es, efs, cgs, fbs = sub[s]
        _, _, elf, aef, rule, _ = fue[f]
        code = f'{f}.{cgs if rule == "sub" else "all"}'
        sp0 = float(prices[f'mit.sp.{code}']) / units[f]
        btax = (float(prices[f'mit.rp.{code}']) - float(prices[f'mit.sp.{code}'])) / units[f]
        e = ef[f'EGY|{f}|{efs}']
        fcov = 0 if f == 'bio' else bool(mt[n['MCov' + f.capitalize()]])
        scov = bool(mt[n['MCov' + s.capitalize()]])
        fbcov = bool(mt[66 + B.COV_SECTORS.index(s)])
        eY, eU, eF = el[f'inc|{elf}|{es}'], el[f'usg|{elf}|{es}'], el[f'eff|{elf}|{es}']
        a = el[f'aei|{aef}|{es}']
        v = {k: [] for k in B.VARS}
        for y in years:
            v['sp'].append(sp0)
            v['ctxnew'].append(cp[y] * e * fcov * scov)
            v['ntx'].append(fpr[code][y] / units[f])
            v['nce'].append(v['ctxnew'][-1] + v['ntx'][-1])
            v['tax'].append(btax + v['nce'][-1])
            v['atp'].append(sp0 + v['tax'][-1])
            v['shp'].append(fb[fbs][y] * e * fbcov)
            if y == B.BASE_YEAR:
                v['ener'].append(use0[(s, f)])
            else:
                ratio = v['atp'][-1] / v['atp'][-2]
                v['ener'].append(v['ener'][-1] * (1 / (1 + a)) ** (1 + eU) * (1 + float(weo[str(y)])) ** eY
                                 * ratio ** eU * ratio ** (eF * (1 + eU)))
        for k in B.VARS:
            res[(k, s, f)] = v[k]
    return res


def scenario_mt(g):
    mt = mt_used()
    for k, v in B.MT_SCENARIO_INPUTS[g].items():
        mt[B.MT_NAMES[k]] = v
    return mt


def compare(ws, g, exp, label):
    _, base, years = B.group_cols(g)
    worst = (0.0, None)
    for var in B.VARS:
        for r, s, f in B.var_rows(var):
            for k, c in enumerate([base] + years):
                v, e = ws.cell(r, c).value, exp[(var, s, f)][k]
                d = float('inf') if not isinstance(v, (int, float)) else abs(v - e) / max(1.0, abs(e))
                if d > worst[0]:
                    worst = (d, f'{L(c)}{r}')
    ok = worst[0] < TOL
    log(f'- {label}: {len(B.VARS)} variables x {len(B.PAIRS)} rows x {B.NY} years, max relative diff '
        f'{worst[0]:.2e} (at {worst[1]}) {"OK" if ok else "FAIL"}')
    return ok


# ------------------------------------------------------------------ checks
def check_shipped(ws):
    lam = {B.group_cols(g)[2][-1] for g, *_ in B.SCENARIOS}
    errs = [(c.row, c.column) for row in ws.iter_rows() for c in row if is_err(c.value)]
    outside = [f'{L(c)}{r}' for r, c in errs if c not in lam]
    ok = not outside
    log(f'- error values: {len(errs)}, all in the {B.YEARS[-1]} (LAMBDA) columns: '
        f'{"OK" if ok else "FAIL, outside: " + str(outside[:10])}')
    return ok


def uniform(ws, cells, label):
    forms = {r1c1(ws.cell(r, c).value, r, c) for r, c in cells}
    ok = len(forms) == 1
    log(f'- {label}: {len(cells)} cells, {len(forms)} formula {"OK" if ok else "FAIL"}')
    return ok, next(iter(forms))


def check_uniformity(wb):
    ws = wb['Mitigation']
    ok = True
    grp = [B.group_cols(g) for g, *_ in B.SCENARIOS]
    for var in B.VARS:
        rows = [r for r, _s, _f in B.var_rows(var)]
        g1, _ = uniform(ws, [(r, b) for r in rows for _, b, ys in grp], f'`{var}` base year')
        g2, f2 = uniform(ws, [(r, c) for r in rows for _, b, ys in grp for c in ys[:-1]],
                         f'`{var}` {B.YEARS[0]}-{B.YEARS[-2]}')
        g3, f3 = uniform(ws, [(r, ys[-1]) for r in rows for _, b, ys in grp], f'`{var}` {B.YEARS[-1]}')
        g4 = all(uniform(ws, [(r, CI(pc)) for r in rows], f'`{var}` column {pc}')[0] for pc in B.PARAM_COLS)
        log(f'  - `{f2[:150]}`')
        if f3 != f2:
            log(f'  - {B.YEARS[-1]}: `{f3[:150]}`')
        ok &= g1 and g2 and g3 and g4
    mt_rows = [B.R_POL[k] for k, _v, _f, _s, kind, _m in B.POL if kind == 'mt']
    ok &= uniform(ws, [(r, c) for r in mt_rows for _, b, ys in grp for c in [b] + ys], 'Section 1 inputs')[0]
    for prefix, n in (('fpr.', len(B.PRICE_FUELS)), ('fb.', len(B.FEEBATE_SECTORS))):
        keys = [k for k, *_ in B.POL if k.startswith(prefix) and k not in ('fpr.yr0', 'fpr.yr1', 'fb.yr0', 'fb.yr1')]
        assert len(keys) == n
        ok &= uniform(ws, [(B.R_POL[k], c) for k in keys for _, b, ys in grp for c in [b] + ys],
                      f'Section 1 paths `{prefix}*`')[0]
    heads = list(B.SUB_HEAD.values())
    ok &= uniform(ws, [(h, c) for h in heads for _, b, ys in grp for c in [b] + ys], 'subsector heading totals')[0]
    hidden = ws.column_dimensions['D'].hidden and ws.column_dimensions['D'].max == 7
    log(f'- parameter columns D:G grouped and hidden: {"OK" if hidden else "FAIL"}')
    return ok and hidden


def check_encoding():
    ok = True
    for name, (params, body) in B.LAMBDAS.items():
        xml = B.lambda_xml(name)
        bare = [p for p in params if re.search(rf'(?<![A-Za-z_.]){p}(?![A-Za-z_\d(])', xml)]
        good = not bare and xml.startswith('_xlfn.LAMBDA(')
        ok &= good
        log(f'- {name}: {len(params)} parameters, {xml.count("_xlpm.")} _xlpm. tokens, bare names {bare} '
            f'{"OK" if good else "FAIL"}')
    return ok


def mt_set(wsm, col, row_or_name, value):
    r = row_or_name if isinstance(row_or_name, int) else B.MT_NAMES[row_or_name]
    wsm.cell(r, col).value = value


def add_scenario(wb, src_g, dst_g, name, edits):
    """Copy MTInputs column of src_g to dst_g's column (keeping the auto number) and the Mitigation group."""
    ws, wsm = wb['Mitigation'], wb['MTInputs']
    ms, md = B.MT_COL0 + src_g - 1, B.MT_COL0 + dst_g - 1
    for r in range(1, 416):
        v = wsm.cell(r, ms).value
        if isinstance(v, str) and v.startswith('='):
            v = Translator(v, origin=f'{L(ms)}{r}').translate_formula(f'{L(md)}{r}')
        wsm.cell(r, md).value = v
    wsm.cell(B.MT_ROW_SCEN, md).value = f'={L(md - 1)}{B.MT_ROW_SCEN}+1'
    wsm.cell(B.MT_ROW_NAME, md).value = name
    for k, v in edits.items():
        mt_set(wsm, md, k, v)
    s0, d0 = B.group_cols(src_g)[0], B.group_cols(dst_g)[0]
    for k in range(B.GW):
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, s0 + k).value
            if isinstance(v, str) and v.startswith('='):
                v = Translator(v, origin=f'{L(s0 + k)}{r}').translate_formula(f'{L(d0 + k)}{r}')
            ws.cell(r, d0 + k).value = v
    ws.cell(B.R_SCEN, d0).value = f'={L(d0 - B.GW)}{B.R_SCEN}+1'


LEGACY_CP = {2022: 0, 2023: 0, 2024: 0, 2025: 0, 2026: 0, 2027: 12.5, 2028: 25, 2029: 37.5, 2030: 50,
             2031: 62.5, 2032: 75, 2033: 87.5, 2034: 100, 2035: 112.5}   # legacy Mitigation row 8582
FPR_TEST = {140: 2027, 141: 2030, 150: 0.05, 164: 0.20, 146: 0.5, 160: 2.0}   # gasoline $/l, gas res $/GJ
FB_TEST = {'D_FeebateIntro': 2027, 'D_FeebateOutro': 2030, 'D_Feb_Level_Start_Trans': 10,
           'D_Feb_Level_Target_Trans': 50, 'D_Feb_Level_Start_Ind': 5, 'D_Feb_Level_Target_Ind': 25,
           67: True, 74: True, 75: True}                                  # feebate coverage: rod, mch, irn


def check_scenarios():
    wb = openpyxl.load_workbook(B.OUT)
    zero = {'CPLevelStart': 0, 'CPLevelTarget': 0}
    add_scenario(wb, 2, 3, 'Test 3: copy of 2', {})
    add_scenario(wb, 3, 4, 'Test 4: copy of 2, carbon price 0', zero)
    add_scenario(wb, 4, 5, 'Test 5: legacy Egypt carbon tax', {'CPIntro': 2026, 'CPLevelStart': 0,
                                                                'CPLevelTarget': 50, 'CPOutro': 2030})
    add_scenario(wb, 4, 6, 'Test 6: fuel price reform', FPR_TEST)
    add_scenario(wb, 4, 7, 'Test 7: feebates', FB_TEST)
    ws = recalc(variant(wb, 'expand'))
    ok = True
    for a, b, label in [(2, 3, 'scenario 3 (copy of 2) = scenario 2'),
                        (1, 4, 'scenario 4 (copy of 2, carbon price 0 in MTInputs) = scenario 1')]:
        ca, cb = B.group_cols(a), B.group_cols(b)
        diff, n = 0.0, 0
        for r in range(B.R_YEAR, ws.max_row + 1):
            if r in (B.R_SCEN, B.R_NAME):
                continue
            for x, y in zip([ca[1]] + ca[2], [cb[1]] + cb[2]):
                va, vb = ws.cell(r, x).value, ws.cell(r, y).value
                if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                    if b == 4 and r in (B.R_POL['CPLevelStart'], B.R_POL['CPLevelTarget']):
                        continue
                    diff, n = max(diff, abs(va - vb)), n + 1
                elif va != vb:
                    diff = float('inf')
        good = diff < TOL
        ok &= good
        log(f'- {label}: {n} numeric cells, max abs diff {diff:.3g} {"OK" if good else "FAIL"}')
    for g in (3, 4, 5, 6, 7):
        code, base, _ = B.group_cols(g)
        num_, name = ws.cell(B.R_SCEN, code).value, ws.cell(B.R_NAME, base).value
        c = ws.cell(B.SUB_HEAD['rod'] + B.VOFF['ener'], code).value
        good = num_ == g and str(c).endswith(f'.e.{g}') and str(name).startswith(f'Test {g}')
        ok &= good
        log(f'- pasted group {g}: number {num_}, name "{name}", code `{c}` {"OK" if good else "FAIL"}')
    _, base, years = B.group_cols(5)
    got = {ws.cell(B.R_YEAR, c).value: ws.cell(B.R_CP, c).value for c in [base] + years}
    d = max(abs(got[y] - v) for y, v in LEGACY_CP.items())
    ok &= d < TOL
    log(f'- scenario 5 (legacy Egypt carbon tax): carbon price {[got[y] for y in (2026, 2027, 2030, 2035)]} vs '
        f'legacy row 8582, max abs diff {d:.3g} {"OK" if d < TOL else "FAIL"}')
    base_mt = scenario_mt(2)
    for g, edits in ((5, {'CPIntro': 2026, 'CPLevelStart': 0, 'CPLevelTarget': 50, 'CPOutro': 2030}),
                     (6, {**zero, **FPR_TEST}), (7, {**zero, **FB_TEST})):
        mt = dict(base_mt)
        for k, v in edits.items():
            mt[k if isinstance(k, int) else B.MT_NAMES[k]] = v
        ok &= compare(ws, g, recompute(mt), f'scenario {g} vs Python')
    r_ntx = B.SUB_HEAD['rod'] + B.VOFF['ntx'] + 2
    r_shp = B.SUB_HEAD['rod'] + B.VOFF['shp'] + 2
    _, b6, y6 = B.group_cols(6)
    _, b7, y7 = B.group_cols(7)
    log(f'  - scenario 6, road gasoline new excise ($/GJ) 2026-2031: '
        f'{[round(ws.cell(r_ntx, c).value, 3) for c in y6[3:9]]}')
    log(f'  - scenario 7, road gasoline feebate shadow price ($/GJ) 2026-2031: '
        f'{[round(ws.cell(r_shp, c).value, 3) for c in y7[3:9]]}')
    (_, b1, y1), (_, b4, y4) = B.group_cols(1), B.group_cols(7)
    d = max(abs(ws.cell(r, x).value - ws.cell(r, y).value) for r, _s, _f in B.var_rows('ener')
            for x, y in zip([b1] + y1, [b4] + y4))
    good = d < TOL and any(abs(ws.cell(r_shp, c).value) > 0 for c in y7)
    ok &= good
    log(f'- scenario 7: shadow price non-zero, fuel use = scenario 1 (shadow price not yet used), max abs diff '
        f'{d:.3g} {"OK" if good else "FAIL"}')
    return ok


def check_regression(ws_new):
    spec = importlib.util.spec_from_file_location('b06', os.path.join(HERE, 'Old', 'build_v0_6.py'))
    b06 = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(b06)
    wbo = openpyxl.load_workbook(os.path.join(HERE, 'Old', 'CPAT_Mitigation_CopyPaste_v0.6.xlsx'))
    wso = wbo['Mitigation']
    for _, (b0, *_r) in b06.BLOCKS.items():
        for r in b06.data_rows(b0):
            for g in (1, 2):
                c = b06.group_cols(g)[2][-1]
                wso.cell(r, c).value = expand(wso.cell(r, c).value, b06.LAMBDAS)
    path = os.path.join(tempfile.mkdtemp(), 'old.xlsx')
    wbo.save(path)
    old = recalc(path)
    vo, vn = {}, {}
    for g in (1, 2):
        for src, mod, dst in ((old, b06, vo), (ws_new, B, vn)):
            code, base, years = mod.group_cols(g)
            for r in range(1, src.max_row + 1):
                k = src.cell(r, code).value
                if isinstance(k, str) and k.startswith('egy.'):
                    dst[k] = [src.cell(r, c).value for c in [base] + years]
    common = sorted(set(vo) & set(vn))
    diff = max(abs(a - b) for k in common for a, b in zip(vo[k], vn[k])
               if isinstance(a, (int, float)) and isinstance(b, (int, float)))
    ok = diff < TOL and len(common) > 1000
    vars_ = sorted({k.split('.')[2] for k in common})
    log(f'- {len(common)} output codes in both versions ({", ".join(vars_)}), max abs diff {diff:.3g} '
        f'{"OK" if ok else "FAIL"}')
    return ok


def check_codes(wsf, wsv):
    ok = True
    described = [r for r in range(1, wsf.max_row + 1)
                 if isinstance(wsf[f'H{r}'].value, str) and wsf[f'H{r}'].value.startswith('=INDEX(Variables')]
    for col in 'HIJ':
        ok &= uniform(wsf, [(r, CI(col)) for r in described], f'column {col}')[0]
    for g, *_ in B.SCENARIOS:
        code = B.group_cols(g)[0]
        ok &= uniform(wsf, [(r, code) for r in described], f'code column {L(code)} (scenario {g})')[0]
    k1, k2 = B.group_cols(1)[0], B.group_cols(2)[0]
    h = B.SUB_HEAD['rod']
    expected = [(h + B.VOFF['ener'] + 2, k1, 'egy.mit.ener.rod.gso.e.1', 'Fuel use | Road | Gasoline'),
                (h + B.VOFF['atp'] + 2, k2, 'egy.mit.atp.rod.gso.e.2', 'After-tax price | Road | Gasoline'),
                (h + B.VOFF['ctxnew'] + 3, k2, 'egy.mit.ctxnew.rod.die.a.2', 'New carbon tax | Road | Diesel'),
                (h + B.VOFF['shp'] + 2, k1, 'egy.mit.shp.rod.gso.1', None),
                (h, k1, 'egy.mit.ener.rod.all.e.1', 'Fuel use | Road | All fuels'),
                (B.SEC_SUM['tra'], k1, 'egy.mit.ener.tra.all.e.1', 'Fuel use | Transport | All fuels'),
                (B.R_CP, k2, 'egy.mit.cptraj.2', 'Carbon price trajectory used'),
                (B.R_POL['sc.oen'], k1, 'egy.mit.ctcov.oen.all.1', None),
                (B.R_TOTAL, k2, 'egy.mit.ener.all.all.e.2', 'Fuel use | All subsectors | All fuels')]
    for r, c, code, label in expected:
        got, lab = wsv.cell(r, c).value, wsv[f'H{r}'].value
        good = got == code and (label is None or lab == label)
        ok &= good
        log(f'- {L(c)}{r}: `{got}` | {lab} {"OK" if good else "FAIL (expected " + code + ")"}')
    return ok


def check_outline(wsf):
    ok = wsf.sheet_properties.outlinePr.summaryBelow is False
    lv = lambda r: (wsf.row_dimensions[r].outlineLevel or 0, bool(wsf.row_dimensions[r].hidden))
    checks = [('carbon price (summary of section 1)', [B.R_CP], (0, False)),
              ('section 1 inputs and paths', [B.R_POL[k] for k, *_ in B.POL if k != 'cptraj'], (1, True)),
              ('sector totals', list(B.SEC_SUM.values()), (0, False)),
              ('subsector headings', list(B.SUB_HEAD.values()), (1, False)),
              ('subsector variables', [r for v in B.VARS for r, *_ in B.var_rows(v)], (2, True)),
              ('results', list(range(B.R_SUB0, B.R_PCT + 1)), (1, False))]
    for label, rows, exp in checks:
        good = all(lv(r) == exp for r in rows)
        ok &= good
        log(f'- {label}: {len(rows)} rows at level {exp[0]}, {"hidden" if exp[1] else "visible"} '
            f'{"OK" if good else "FAIL"}')
    return ok


def check_mtinputs():
    tpl = openpyxl.load_workbook(B.MT_TEMPLATE)['MTInputs']
    ws = openpyxl.load_workbook(B.OUT)['MTInputs']
    diffs = [f'{L(c)}{r}' for r in range(1, tpl.max_row + 1) for c in range(1, 9)
             if tpl.cell(r, c).value != ws.cell(r, c).value]
    ok = not diffs and ws.column_dimensions['D'].hidden and ws.column_dimensions['E'].hidden
    log(f'- columns A:H, rows 1-{tpl.max_row}: {len(diffs)} cells differ from the template; D:E hidden '
        f'{"OK" if ok else "FAIL " + str(diffs[:10])}')
    return ok


def check_format(wsf):
    bands = [3, B.B_POL, B.B_POW, B.B_RES] + list(B.SEC_BAND.values())
    ok = all(str(wsf[f'A{r}'].font.color.rgb).endswith('FFFFFF') and wsf[f'A{r}'].font.b for r in bands)
    texts = [wsf[f'A{r}'].value[:22] for r in sorted(bands)]
    log(f'- section bands {texts}: white bold text {"OK" if ok else "FAIL"}')
    return ok


def check_sanity(ws):
    ok = True
    (_, b1, y1), (_, b2, y2) = B.group_cols(1), B.group_cols(2)
    cols1, cols2, years = [b1] + y1, [b2] + y2, [B.BASE_YEAR] + B.YEARS
    pre = max(abs(ws.cell(r, x).value - ws.cell(r, y).value) for r, *_ in B.var_rows('ener')
              for x, y, yr in zip(cols1, cols2, years) if yr < 2027)
    ok &= pre < TOL
    log(f'- scenario 1 vs 2 before 2027 (fuel use): max abs diff {pre:.3g} {"OK" if pre < TOL else "FAIL"}')
    oen = max(abs(ws.cell(r, x).value - ws.cell(r, y).value) for r, s, f in B.var_rows('ener') if s == 'oen'
              for x, y in zip(cols1, cols2))
    ok &= oen < TOL
    log(f'- other energy use not taxed (MCovOen FALSE): fuel use scenario 1 = 2, max abs diff {oen:.3g} '
        f'{"OK" if oen < TOL else "FAIL"}')
    chk = max(abs(ws.cell(B.R_CHK, c).value) for c in cols1 + cols2)
    secs = max(abs(sum(ws.cell(B.SEC_SUM[g], c).value for g in B.SEC_SUM) - ws.cell(B.R_TOTAL, c).value)
               for c in cols1 + cols2)
    ok &= chk < 1e-6 and secs < 1e-6
    log(f'- check row {B.R_CHK}: max abs {chk:.3g}; sector totals sum to total: max abs {secs:.3g} '
        f'{"OK" if chk < 1e-6 and secs < 1e-6 else "FAIL"}')
    log('')
    log('| Year | Scenario 1 (ktoe) | Scenario 2 (ktoe) | Change |')
    log('|---|---|---|---|')
    for x, y, yr in zip(cols1, cols2, years):
        if yr in (2022, 2026, 2027, 2030, 2035):
            a, b = ws.cell(B.R_TOTAL, x).value, ws.cell(B.R_TOTAL, y).value
            log(f'| {yr} | {a:,.0f} | {b:,.0f} | {b / a - 1:+.1%} |')
    return ok


def main():
    log(f'# Check report - CPAT_Mitigation_CopyPaste_v{B.VERSION}')
    results = []
    log('\n## 1. Shipped file recalculated in LibreOffice (no LAMBDA support)')
    results.append(check_shipped(recalc(B.OUT)))
    log('\n## 2. Formula uniformity (relative R1C1)')
    results.append(check_uniformity(openpyxl.load_workbook(B.OUT)))
    log('\n## 3. LAMBDA encoding')
    results.append(check_encoding())
    log('\n## 4. Values vs independent Python recomputation')
    ws_a = recalc(variant(openpyxl.load_workbook(B.OUT), 'expand'))
    ws_b = recalc(variant(openpyxl.load_workbook(B.OUT), 'drag'))
    for g, *_ in B.SCENARIOS:
        exp = recompute(scenario_mt(g))
        results.append(compare(ws_a, g, exp, f'A. LAMBDA expanded, scenario {g}'))
        results.append(compare(ws_b, g, exp, f'B. plain formula dragged, scenario {g}'))
    log('\n## 5. Scenario tests (copied MTInputs column + Mitigation group)')
    results.append(check_scenarios())
    log('\n## 6. Regression vs v0.6 (shared output codes)')
    results.append(check_regression(ws_a))
    log('\n## 7. Labels and codes')
    results.append(check_codes(openpyxl.load_workbook(B.OUT)['Mitigation'], ws_a))
    log('\n## 8. Row outline')
    results.append(check_outline(openpyxl.load_workbook(B.OUT)['Mitigation']))
    log('\n## 9. MTInputs')
    results.append(check_mtinputs())
    log('\n## 10. Format')
    results.append(check_format(openpyxl.load_workbook(B.OUT)['Mitigation']))
    log('\n## 11. Sanity and results')
    results.append(check_sanity(ws_a))
    log(f'\n**Overall: {"PASS" if all(results) else "FAIL"}**')
    with open(os.path.join(HERE, f'check_report_v{B.VERSION}.md'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(report) + '\n')


if __name__ == '__main__':
    main()
