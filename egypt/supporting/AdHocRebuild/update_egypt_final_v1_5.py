"""Update the plain hand-over folder egypt-final/ to the v1.5 numbers (3B rebate to all covered industry).
Edits the three Word files in place (python-docx, untracked) and replaces Egypt_Model.xlsx with kernel v1.5. The CBAM
obligations workbook is retired: its content is the kernel sheet Table2_Final. Run once; it asserts the old values."""
import copy
import os
import shutil

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
D = os.path.join(ROOT, "egypt-final")


def set_text(par, text):
    rs = par.runs
    rs[0].text = text
    for r in rs[1:]:
        r._element.getparent().remove(r._element)


def replace_in(par, old, new):
    full = "".join(r.text for r in par.runs)
    i = full.find(old)
    if i < 0:
        return False
    pos, first = 0, True
    for r in par.runs:
        t = r.text
        a, b = max(i, pos), min(i + len(old), pos + len(t))
        if a < b:
            r.text = (t[:a - pos] + new + t[b - pos:]) if first else (t[:a - pos] + t[b - pos:])
            first = False
        pos += len(t)
    return True


def all_paras(d):
    return [Paragraph(p, d) for p in d.element.body.iter(qn("w:p"))]


def sub(d, old, new, startswith=None):
    for p in all_paras(d):
        if (startswith is None or p.text.startswith(startswith)) and old in p.text:
            assert replace_in(p, old, new)
            return
    raise AssertionError(("not found", old[:60]))


def row_of(d, label, nth=0):
    hits = [r for t in d.tables for r in t.rows if r.cells[0].text.strip() == label]
    assert len(hits) > nth, ("row not found", label)
    return hits[nth]


def set_cell(row, col, old, new):
    c = row.cells[col]
    assert c.text.strip() == old, (row.cells[0].text, c.text, old)
    set_text(c.paragraphs[0], new)


# ---- 1_Summary
f = os.path.join(D, "1_Summary.docx")
d = docx.Document(f)
r3b = next(r for t in d.tables for r in t.rows if r.cells[0].text.strip() == "3B")
c = r3b.cells[1]
assert "CBAM producers get free allowances" in c.text
set_text(c.paragraphs[0], "As 3A, but all covered industry gets free allowances per unit of output")
for label, old, new in (("Emissions cut (Mt CO₂e)", "−19.1", "−12.0"),
                        ("Emissions cut (% of national total)", "−3.2", "−2.0"),
                        ("Carbon revenue (USD bn)", "0.6", "0.3"),
                        ("Air-pollution deaths avoided", "491", "330")):
    set_cell(row_of(d, label), 5, old, new)
sub(d, "Pricing industry only (3A–3B) cuts less (19–23 Mt) and raises much less, because only part of industry was priced in the model run.",
    "Pricing industry only (3A–3B) cuts less (12–23 Mt) and raises much less, because only part of industry was priced in the model run. "
    "In 3B the free allowances cover all covered industry, so net revenue is only about USD 0.3bn and the cut is the smallest.",
    startswith="Pricing industry only")
d.save(f)

# ---- 2_Results_Table
f = os.path.join(D, "2_Results_Table.docx")
d = docx.Document(f)
sub(d, "(sheet CarveOut_Table2)", "(sheets CarveOut_Table2 and Table2_Final)", startswith="Results come from")
for label, old, new, nth in (("Carbon revenue (USD bn)", "0.6", "0.3", 0),
                             ("Emissions cut (Mt CO₂e)", "−19.1", "−12.0", 0),
                             ("Emissions cut (% of national total)", "−3.2", "−2.0", 0),
                             ("Deaths avoided", "491", "330", 0),
                             ("Emissions cut (Mt): final", "−19.1", "−12.0", 0),
                             ("Revenue (USD bn): final", "0.6", "0.3", 0),
                             ("Deaths avoided: final", "491", "330", 0),
                             ("Final cut", "−19.1", "−12.0", 0)):
    set_cell(row_of(d, label, nth), 5, old, new)
add = row_of(d, "Add the separate model’s cut in CBAM goods")
new_row = copy.deepcopy(add._tr)
add._tr.addnext(new_row)
from docx.table import _Row
nr = _Row(new_row, add._parent)
set_text(nr.cells[0].paragraphs[0], "Add back: 3B free allowances for the rest of covered industry")
for j in range(1, 7):
    set_text(nr.cells[j].paragraphs[0], "+7.2" if j == 5 else "0.0")
sub(d, "less the 3B rebate and the 3C fund.", "less the 3B rebate and the 3C fund. In 3B the rebate covers all covered industry, not only the CBAM goods.",
    startswith="Revenue is CPAT")
sub(d, "The CBAM-bill row (Egypt_CBAM_Obligations.xlsx) is a working estimate awaiting confirmation.",
    "The CBAM-bill row (calculated in Egypt_Model.xlsx, sheet Table2_Final) is a working estimate awaiting confirmation.",
    startswith="The CBAM-bill row")
anchor = next(p for p in all_paras(d) if p.text.startswith("Deaths avoided are CPAT"))
el = copy.deepcopy(anchor._element)
anchor._element.addnext(el)
p = Paragraph(el, anchor._parent)
set_text(p, "3B free allowances for all covered industry: in the rest of covered industry the lower-output part (two thirds) of CPAT’s "
            "response is removed, which reduces the cut by 7.2 Mt, and the value of the allowances (USD 0.33bn) comes off revenue. "
            "Rows may not add exactly because of rounding.")
d.save(f)

# ---- 3_Methodology
f = os.path.join(D, "3_Methodology.docx")
d = docx.Document(f)
r3b = next(r for t in d.tables for r in t.rows if r.cells[0].text.strip() == "3B")
c = next(c for c in r3b.cells if c.text.strip() == "As 3A, with free allowances per unit of output")
set_text(c.paragraphs[0], "As 3A, with free allowances per unit of output for all covered industry")
sub(d, "Calculation: Egypt_CBAM_Obligations.xlsx.", "Calculation: Egypt_Model.xlsx, sheet Table2_Final.")
sub(d, "A version with that credit is kept in the workbook for reference only.",
    "A version with that credit is kept in the workbook for reference only, on the 2030 phase-in and with no phase-in.",
    startswith="We take no credit")
anchor = next(p for p in all_paras(d) if p.text.startswith("Baseline output grows"))
el = copy.deepcopy(anchor._element)
anchor._element.addnext(el)
p = Paragraph(el, anchor._parent)
set_text(p, "3B free allowances. In 3B the free allowances cover all covered industry, not only the CBAM goods. In the rest of covered "
            "industry, the lower-output part (one minus the efficiency share, two thirds) of CPAT’s response, including its proportional "
            "process emissions, is removed, and the value of the allowances (carbon price × priced industrial fuel emissions outside the "
            "CBAM goods, USD 0.33bn) is deducted from revenue. Deaths avoided are adjusted for the fuel part. The result is a smaller cut "
            "(−12.0 Mt) and net revenue of USD 0.3bn. Receipts are not recalculated for the changed emissions.")
d.save(f)

# ---- workbooks
shutil.copyfile(os.path.join(ROOT, "egypt", "final", "CPAT_Industry_Kernel_Egypt_v1.5.xlsx"), os.path.join(D, "Egypt_Model.xlsx"))
ob = os.path.join(D, "Egypt_CBAM_Obligations.xlsx")
if os.path.exists(ob):
    os.remove(ob)
print("updated egypt-final")
