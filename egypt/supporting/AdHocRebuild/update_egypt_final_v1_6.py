"""Update egypt-final/ (plain hand-over folder) to v1.6 numbers (product-specific output elasticities).
Usage: python update_egypt_final_v1_6.py [results.json]   (default: carveout_v1_6_results.json, the Python mirror;
use carveout_v1_6_results_live.json after build_v1_6.py). Edits the three Word files in place."""
import json
import os
import sys

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
D = os.path.join(ROOT, "egypt-final")
B = ["1A", "2A", "2B", "3A", "3B", "3C"]
M_CBAM = {"1A": 100, "2A": 44, "2B": 44, "3A": 100, "3B": 100, "3C": 100}
MINUS = "−"
arg = sys.argv[1] if len(sys.argv) > 1 else "carveout_v1_6_results.json"
raw = json.load(open(os.path.join(HERE, arg), encoding="utf8"))
R = raw["results"] if "results" in raw else raw
OLD = json.load(open(os.path.join(HERE, "carveout_v1_5_results.json"), encoding="utf8"))
for b in B:
    R[b].setdefault("O_final", OLD[b]["O_final"])
    R[b]["M"] = M_CBAM[b]


def rnd(x, dp):
    q = 10 ** dp
    return (int(abs(x) * q + 0.5 + 1e-9) / q) * (1 if x >= 0 else -1)


def fmt(x, dp=1, plus=False, comma=False):
    v = rnd(x, dp)
    s = (("{:,.%df}" if comma else "{:.%df}") % dp).format(abs(v))
    if v < 0:
        return MINUS + s
    return ("+" + s) if plus and v > 0 else s


def set_text(par, text):
    rs = par.runs
    rs[0].text = text
    for r in rs[1:]:
        r._element.getparent().remove(r._element)


def replace_in(par, old, new):
    full = "".join(r.text for r in par.runs)
    i = full.find(old)
    if i < 0:
        return False
    pos, first = 0, True
    for r in par.runs:
        t = r.text
        a, b = max(i, pos), min(i + len(old), pos + len(t))
        if a < b:
            r.text = (t[:a - pos] + new + t[b - pos:]) if first else (t[:a - pos] + t[b - pos:])
            first = False
        pos += len(t)
    return True


def sub(d, old, new):
    for p in [Paragraph(x, d) for x in d.element.body.iter(qn("w:p"))]:
        if old in p.text:
            assert replace_in(p, old, new), old
            return
    raise AssertionError(("not found", old[:60]))


def put_row(d, label, vals):
    for t in d.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == label:
                for j, v in enumerate(vals):
                    set_text(r.cells[1 + j].paragraphs[0], v)
                return
    raise AssertionError(("row not found", label))


def main():
    ks = [-R[b]["K"] for b in ("1A", "2A", "2B")]
    lo, hi = rnd(min(ks), 0), rnd(max(ks), 0)
    pl = [-R[b]["L"] for b in ("1A", "2A", "2B")]
    # ---- summary
    f = os.path.join(D, "1_Summary.docx")
    d = docx.Document(f)
    put_row(d, "Emissions cut (Mt CO₂e)", [fmt(R[b]["K"]) for b in B])
    put_row(d, "Emissions cut (% of national total)", [fmt(R[b]["L"]) for b in B])
    put_row(d, "Carbon revenue (USD bn)", [fmt(R[b]["P"]) for b in B])
    put_row(d, "Air-pollution deaths avoided", [fmt(R[b]["Q"], 0, comma=True) for b in B])
    put_row(d, "CBAM goods: emissions per tonne (%)", [fmt(R[b]["N"]) for b in B])
    put_row(d, "CBAM bill per tonne exported (%)", [fmt(R[b]["O_final"]) for b in B])
    sub(d, "cuts emissions by about 33–37 Mt, around 6% of national emissions", "cuts emissions by about %d–%d Mt, around %d–%d%% of national emissions" % (lo, hi, rnd(min(pl), 0), rnd(max(pl), 0)))
    sub(d, "cuts less (12–23 Mt)", "cuts less (%d–%d Mt)" % (rnd(-R["3B"]["K"], 0), rnd(-R["3A"]["K"], 0)))
    sub(d, "Output falls by 0.5% for every 1% rise in net cost. This is a placeholder.",
        "Output falls by 0.1% (cement) to 0.5% (aluminium) for every 1% rise in net cost. These are judgements from the literature; only cement matters.")
    d.save(f)
    # ---- results table
    f = os.path.join(D, "2_Results_Table.docx")
    d = docx.Document(f)
    put_row(d, "Share of national emissions priced (%)", [fmt(R[b]["J"], 0) for b in B])
    put_row(d, "Carbon revenue (USD bn)", [fmt(R[b]["P"]) for b in B])
    put_row(d, "Emissions cut (Mt CO₂e)", [fmt(R[b]["K"]) for b in B])
    put_row(d, "Emissions cut (% of national total)", [fmt(R[b]["L"]) for b in B])
    put_row(d, "CBAM goods: emissions per tonne (%)", [fmt(R[b]["N"]) for b in B])
    put_row(d, "CBAM bill per tonne exported (%)", [fmt(R[b]["O_final"]) for b in B])
    put_row(d, "CBAM goods: total emissions (%)", [fmt(R[b]["T"]) for b in B])
    put_row(d, "Deaths avoided", [fmt(R[b]["Q"], 0, comma=True) for b in B])
    put_row(d, "Emissions cut (Mt): final", [fmt(R[b]["K"]) for b in B])
    put_row(d, "Revenue (USD bn): final", [fmt(R[b]["P"]) for b in B])
    put_row(d, "Deaths avoided: final", [fmt(R[b]["Q"], 0, comma=True) for b in B])
    put_row(d, "CBAM bill (%): final", [fmt(R[b]["O_final"]) for b in B])
    put_row(d, "CPAT total cut", [fmt(R[b]["dghg"]) for b in B])
    put_row(d, "Remove CPAT’s cut in CBAM goods", [fmt(-R[b]["cpat_blk"], 1, plus=True) for b in B])
    put_row(d, "Add the separate model’s cut in CBAM goods", [fmt(R[b]["dB"], 1, plus=True) for b in B])
    put_row(d, "Add back: 3B free allowances for the rest of covered industry",
            [fmt(R[b]["D_nb"], 1, plus=True) if abs(R[b]["D_nb"]) > 0.05 else "0.0" for b in B])
    put_row(d, "Final cut", [fmt(R[b]["K"]) for b in B])
    sub(d, "The CBAM goods emit", "The CBAM goods emit")
    d.save(f)
    # ---- methodology
    f = os.path.join(D, "3_Methodology.docx")
    d = docx.Document(f)
    sub(d, "Output = baseline output × (1 + net cost increase)^(−0.5).",
        "Output = baseline output × (1 + net cost increase)^ε, with ε = −0.10 for cement, −0.40 for steel and fertilisers and −0.50 for aluminium.")
    for t in d.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == "Output response":
                set_text(r.cells[1].paragraphs[0], "−0.1 to −0.5")
                set_text(r.cells[2].paragraphs[0], "Literature: cement industry demand −0.02 to −0.16 with 20–40% pass-through; steel pass-through 55–85%; "
                                                     "world-priced goods higher. Only cement matters")
                set_text(r.cells[3].paragraphs[0], "Low–Medium")
    d.save(f)
    print("updated egypt-final to v1.6 numbers (%s)" % arg)


main()
