"""
Shared paths, table layout specs, and I/O helpers for the cpat_excel data
pipeline.

Forward direction (lossy, each stage simplifies the previous one):
    original (.xlsb) -> data_bymodule -> data_tabular -> data_standardized

Reverse direction (rebuilds the intermediate tiers from data_standardized,
to check that the standardized form still carries what the tabular/bymodule
tiers need):
    data_standardized -> data_tabular_regenerated -> data_bymodule_regenerated

A table's shape going forward (raw grid -> single table -> long/tidy table)
is described once, in MODULES below, and reused by every script so the two
directions can't drift apart.
"""
import os
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

CPAT_EXCEL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_XLSB = os.path.join(
    CPAT_EXCEL_DIR, "original", "CPAT 1.0pre_456_NoPropData.xlsb"
)

DIR_BYMODULE = os.path.join(CPAT_EXCEL_DIR, "data_bymodule")
DIR_TABULAR = os.path.join(CPAT_EXCEL_DIR, "data_tabular")
DIR_STANDARDIZED = os.path.join(CPAT_EXCEL_DIR, "data_standardized")
DIR_TABULAR_REGEN = os.path.join(CPAT_EXCEL_DIR, "data_tabular_regenerated")
DIR_BYMODULE_REGEN = os.path.join(CPAT_EXCEL_DIR, "data_bymodule_regenerated")

BASE_FONT = Font(name="Arial", size=10)
HEADER_FONT = Font(name="Arial", size=10, bold=True)

# BIFF12 error-code byte -> Excel error literal. pyxlsb surfaces these as hex
# strings (e.g. '0x2a'); untranslated they leak into cells as gibberish.
XLSB_ERROR_CODES = {
    "0x0": "#NULL!",
    "0x7": "#DIV/0!",
    "0xf": "#VALUE!",
    "0x17": "#REF!",
    "0x1d": "#NAME?",
    "0x24": "#NUM!",
    "0x2a": "#N/A",
    "0x2b": "#GETTING_DATA",
}


def clean_value(v):
    if isinstance(v, str) and v.startswith("0x"):
        return XLSB_ERROR_CODES.get(v, v)
    return v


# ---------------------------------------------------------------------------
# Module registry: for each module, which source sheets it draws from, how
# each sheet decomposes into clean tables (data_tabular), and how each table
# decomposes into id/variable/value columns for the long form
# (data_standardized). Extend MODULES with new entries (Mitigation, Air
# pollution, Transport, ...) to run the same pipeline for another module.
# ---------------------------------------------------------------------------

# Mapping sheet layout: 0-indexed [header_row, data_start, data_end, cols]
# plus the original title text/position, so data_bymodule can be rebuilt.
_MAPPING_TABLES = [
    {
        "table": "Mapping_SectorCrosswalk",
        "title": "Sectoral mapping (IEA, GTAP, and EORA to CPAT sectors)",
        "title_row": 0,
        "header_row": 3,
        "data_start": 4,
        "data_end": 165,
        "cols": list(range(1, 20)),  # B:T
        "ffill": [],
        "dedupe": False,
    },
    {
        "table": "Mapping_CPATSectorsToISIC",
        "title": "CPAT sectors to ISIC",
        "title_row": 169,
        "header_row": 172,
        "data_start": 173,
        "data_end": 201,
        "cols": [2, 3, 4, 5, 6],  # C:G
        "ffill": [0, 1, 2],  # forward-fill merged main/sub sector labels
        "dedupe": True,  # rows are duplicated across the paired IEA table
    },
    {
        "table": "Mapping_IEAFlowsToISIC",
        "title": None,  # shares the "CPAT sectors to ISIC" title/row block
        "title_row": None,
        "header_row": 172,
        "data_start": 173,
        "data_end": 201,
        "cols": [8, 9, 10],  # I:K
        "ffill": [],
        "dedupe": False,
    },
    {
        "table": "Mapping_ISICToCPAT",
        "title": None,
        "title_row": None,
        "header_row": 203,
        "data_start": 204,
        "data_end": 316,
        "cols": [3, 4, 5, 6],  # D:G
        "ffill": [],
        "dedupe": False,
    },
    {
        "table": "Mapping_CountriesToGTAP10",
        "title": "CPAT countries to GTAP-10 country groups",
        "title_row": 320,
        "header_row": 323,
        "data_start": 324,
        "data_end": 529,
        "cols": [1, 2, 3],  # B:D
        "ffill": [],
        "dedupe": False,
    },
]

MODULES = {
    "Distribution": {
        "output_name": "CPAT_DistributionalData.xlsx",
        "source_sheets": [
            "Mapping",
            "IO_GTAP",
            "HHSurvey",
            "HH_Elast",
            "ASPIRE",
            "WHOCooking",
            "GDPRatios",
        ],
        # sheets that are already a single clean table: header on row 0,
        # data below, every column kept -> identical in bymodule & tabular
        "simple_sheets": [
            "IO_GTAP",
            "HHSurvey",
            "HH_Elast",
            "ASPIRE",
            "WHOCooking",
            "GDPRatios",
        ],
        # sheets that bundle several logical tables (handled specially)
        "composite_sheets": {"Mapping": _MAPPING_TABLES},
        # data_tabular table name -> the data_bymodule sheet it came from
        "table_to_sheet": {
            "IO_GTAP": "IO_GTAP",
            "HHSurvey": "HHSurvey",
            "HH_Elast": "HH_Elast",
            "ASPIRE": "ASPIRE",
            "WHOCooking": "WHOCooking",
            "GDPRatios": "GDPRatios",
            "Mapping_SectorCrosswalk": "Mapping",
            "Mapping_CPATSectorsToISIC": "Mapping",
            "Mapping_IEAFlowsToISIC": "Mapping",
            "Mapping_ISICToCPAT": "Mapping",
            "Mapping_CountriesToGTAP10": "Mapping",
        },
        # data_tabular table name -> melt spec (None = already tidy, kept as-is)
        "melt_specs": {
            "IO_GTAP": {
                "id_cols": ["dis", "iso3", "year", "sector"],
                "var_name": "gtap_sector",
                "value_name": "value",
            },
            "HHSurvey": {
                "id_cols": ["code", "iso3", "year", "sample", "type", "stat_type", "quant_cons"],
                "var_name": "variable",
                "value_name": "value",
            },
            "HH_Elast": {
                "id_cols": ["code", "iso3", "year", "sample", "type", "stat_type", "quant_cons"],
                "var_name": "variable",
                "value_name": "value",
            },
            "ASPIRE": None,  # already one row per indicator/value
            "WHOCooking": {
                "id_cols": ["code", "iso3", "year", "sample"],
                "var_name": "variable",
                "value_name": "value",
            },
            "GDPRatios": {
                "id_cols": ["Country Name", "Country Code", "Indicator Name", "Indicator Code", "Latest Year"],
                "var_name": "year",
                "value_name": "value",
            },
            "Mapping_SectorCrosswalk": None,
            "Mapping_CPATSectorsToISIC": None,
            "Mapping_IEAFlowsToISIC": None,
            "Mapping_ISICToCPAT": None,
            "Mapping_CountriesToGTAP10": None,
        },
    }
}


# ---------------------------------------------------------------------------
# Grid / table helpers
# ---------------------------------------------------------------------------

def subgrid(grid, header_row, data_start, data_end, cols):
    """Extract a header row + data rows for the given 0-indexed cols."""
    header = [grid[header_row][c] for c in cols]
    data = [[grid[r][c] for c in cols] for r in range(data_start, data_end + 1)]
    return header, data


def ffill_cols(header, data, col_positions):
    """Forward-fill blanks in-place for the given 0-indexed column positions."""
    last = {p: None for p in col_positions}
    for row in data:
        for p in col_positions:
            if row[p] is None:
                row[p] = last[p]
            else:
                last[p] = row[p]
    return header, data


def dedupe_rows(header, data):
    seen = set()
    out = []
    for row in data:
        key = tuple(row)
        if key not in seen:
            seen.add(key)
            out.append(row)
    return header, out


def melt_rows(header, data, id_cols, var_name, value_name):
    """Long-format a wide table: id_cols stay fixed, every other column
    becomes one (variable, value) row."""
    id_idx = [header.index(c) for c in id_cols]
    value_idx = [i for i in range(len(header)) if i not in id_idx]
    out_header = [header[i] for i in id_idx] + [var_name, value_name]
    out_data = []
    for row in data:
        base = [row[i] for i in id_idx]
        for vi in value_idx:
            out_data.append(base + [header[vi], row[vi]])
    return out_header, out_data


def pivot_rows(header, data, id_cols, var_name, value_name):
    """Inverse of melt_rows: rebuild the wide table from a long one, in the
    original id-then-variable row/column order."""
    id_idx = [header.index(c) for c in id_cols]
    var_idx = header.index(var_name)
    val_idx = header.index(value_name)

    columns = []  # variable values, first-seen order
    seen_cols = set()
    groups = OrderedDict()  # id tuple -> {variable: value}
    for row in data:
        key = tuple(row[i] for i in id_idx)
        var = row[var_idx]
        if var not in seen_cols:
            seen_cols.add(var)
            columns.append(var)
        groups.setdefault(key, {})[var] = row[val_idx]

    out_header = list(id_cols) + columns
    out_data = []
    for key, values in groups.items():
        out_data.append(list(key) + [values.get(c) for c in columns])
    return out_header, out_data


# ---------------------------------------------------------------------------
# Workbook I/O
# ---------------------------------------------------------------------------

def read_tables_xlsx(path):
    """Read every sheet of an xlsx workbook as {sheet_name: (header, data)}."""
    wb = load_workbook(path, read_only=True, data_only=True)
    tables = OrderedDict()
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        header = list(rows[0]) if rows else []
        data = [list(r) for r in rows[1:]]
        tables[name] = (header, data)
    wb.close()
    return tables


def write_tables_xlsx(path, tables, header_row=True, freeze=True):
    """Write {sheet_name: (header, data)} to an xlsx workbook, one sheet
    per table, with light formatting."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, data) in tables.items():
        ws = wb.create_sheet(title=name[:31])
        start_row = 1
        if header_row and header:
            for c, val in enumerate(header, start=1):
                cell = ws.cell(row=1, column=c, value=val)
                cell.font = HEADER_FONT
            start_row = 2
        for r, row in enumerate(data, start=start_row):
            for c, val in enumerate(row, start=1):
                if val is None:
                    continue
                cell = ws.cell(row=r, column=c, value=val)
                cell.font = BASE_FONT
        ncols = max(len(header), max((len(r) for r in data), default=0))
        for c in range(1, ncols + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16
        if header_row and header and freeze:
            ws.freeze_panes = "A2"
            if data:
                ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{len(data) + 1}"
    wb.save(path)
    return path


def write_raw_grids_xlsx(path, grids):
    """Write {sheet_name: grid} (list-of-lists, values placed at their exact
    row/col position, no header styling) to an xlsx workbook."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, grid in grids.items():
        ws = wb.create_sheet(title=name[:31])
        for r, row in enumerate(grid, start=1):
            for c, val in enumerate(row, start=1):
                if val is None:
                    continue
                cell = ws.cell(row=r, column=c, value=val)
                cell.font = BASE_FONT
        ncols = max((len(row) for row in grid), default=0)
        for c in range(1, ncols + 1):
            ws.column_dimensions[get_column_letter(c)].width = 14
    wb.save(path)
    return path


def compare_tables(expected, actual, label=""):
    """Print a short diff summary between two {name: (header, data)} dicts.
    Returns True if they match exactly."""
    ok = True
    names = sorted(set(expected) | set(actual))
    for name in names:
        if name not in expected:
            print(f"  [{label}] {name}: unexpected extra table")
            ok = False
            continue
        if name not in actual:
            print(f"  [{label}] {name}: missing table")
            ok = False
            continue
        eh, ed = expected[name]
        ah, ad = actual[name]
        if eh != ah:
            print(f"  [{label}] {name}: header mismatch\n    expected {eh}\n    actual   {ah}")
            ok = False
            continue
        if len(ed) != len(ad):
            print(f"  [{label}] {name}: row count mismatch (expected {len(ed)}, got {len(ad)})")
            ok = False
            continue
        mismatches = sum(1 for e, a in zip(ed, ad) if e != a)
        if mismatches:
            print(f"  [{label}] {name}: {mismatches}/{len(ed)} data rows differ")
            ok = False
    if ok:
        print(f"  [{label}] all {len(names)} tables match exactly")
    return ok
